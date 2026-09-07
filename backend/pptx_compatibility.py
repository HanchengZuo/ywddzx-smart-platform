"""Office 2007-compatible, editable evidence for every report renderer.

Charts keep their embedded workbooks. Visible values use real table cells rather
than c:dTable, whose rendering (especially zeroes) varies between office suites.
"""
import math
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_TICK_LABEL_POSITION
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

COMPATIBILITY_VERSION = 'editable-ooxml-v6.7-1'
DATA_TABLE_PREFIX = 'report-compatible-values-'
FONT = 'Microsoft YaHei'
BAR_TYPES = {XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.BAR_CLUSTERED,
             XL_CHART_TYPE.COLUMN_STACKED, XL_CHART_TYPE.BAR_STACKED}


def metric_text(value):
    if value is None:
        return '—'
    number = float(value)
    if not math.isfinite(number):
        raise ValueError('图表包含非有限数值。')
    return format(number, '.8f').rstrip('0').rstrip('.') if number else '0'


def _font_properties(properties, font_name=FONT):
    latin = properties.find('{http://schemas.openxmlformats.org/drawingml/2006/main}latin')
    name = latin.get('typeface') if latin is not None else font_name
    if not name or name.startswith('+'):
        name = font_name
    for tag in ('a:latin','a:ea','a:cs'):
        node = properties.find('{http://schemas.openxmlformats.org/drawingml/2006/main}' + tag.split(':')[1])
        if node is None:
            node = OxmlElement(tag)
            following = {'a:latin': ('ea','cs','sym','hlinkClick','hlinkMouseOver','extLst'),
                         'a:ea': ('cs','sym','hlinkClick','hlinkMouseOver','extLst'),
                         'a:cs': ('sym','hlinkClick','hlinkMouseOver','extLst')}[tag]
            next_child = next((child for child in properties if child.tag.rsplit('}',1)[-1] in following),None)
            if next_child is not None:
                properties.insert(properties.index(next_child),node)
            else:
                properties.append(node)
        node.set('typeface', name)


def _add_overflow_tables(prs, slide, shape, categories, series):
    """Keep dense evidence readable on adjacent data pages instead of clipping it."""
    shape.name += '-compatible-continuation'
    for node in shape.chart._chartSpace.xpath('.//c:dTable'):
        node.getparent().remove(node)
    shape.height -= Pt(18)
    note = slide.shapes.add_textbox(shape.left, shape.top + shape.height, shape.width, Pt(18))
    note.text_frame.text = '完整数值（含0值）见后续图表数据明细页'
    note.text_frame.paragraphs[0].font.size = Pt(9)
    ids = prs.slides._sldIdLst
    insert_at = list(prs.slides).index(slide) + 1
    for cat_start in range(0, len(categories), 10):
        for series_start in range(0, len(series), 8):
            page = prs.slides.add_slide(prs.slide_layouts[-1])
            for placeholder in list(page.shapes):
                placeholder._element.getparent().remove(placeholder._element)
            title = page.shapes.add_textbox(Inches(.5), Inches(.3), prs.slide_width - Inches(1), Inches(.6))
            title.text_frame.text = '图表数据明细（续）'
            title.text_frame.paragraphs[0].font.size = Pt(22)
            names = categories[cat_start:cat_start + 10]
            selected = series[series_start:series_start + 8]
            rows = [['指标'] + names] + [[s.name or '问题数量'] + [metric_text(v) for v in s.values[cat_start:cat_start + 10]] for s in selected]
            table = page.shapes.add_table(len(rows), len(names) + 1, Inches(.5), Inches(1.1),
                prs.slide_width - Inches(1), prs.slide_height - Inches(1.6)).table
            for r, values in enumerate(rows):
                for c, value in enumerate(values):
                    cell = table.cell(r, c); cell.text = value
                    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                    cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string('EDF2F7' if r == 0 else 'FFFFFF')
                    for p in cell.text_frame.paragraphs:
                        p.alignment = PP_ALIGN.CENTER
                        for run in p.runs:
                            run.font.size = Pt(11); run.font.name = FONT
                            run.font.color.rgb = RGBColor.from_string('334155')
                            _font_properties(run._r.get_or_add_rPr())
            element = ids[-1]; ids.remove(element); ids.insert(insert_at, element)
            insert_at += 1


def add_editable_value_table(prs, slide, shape):
    chart = shape.chart
    if chart.chart_type not in BAR_TYPES or shape.name.endswith('-compatible-continuation') or any(s.name == DATA_TABLE_PREFIX + str(shape.shape_id) for s in slide.shapes):
        return
    categories = [str(c.label) for c in chart.plots[0].categories]
    series = list(chart.series)
    if not categories or not series:
        return
    width, height = shape.width, shape.height
    first_width = min(Inches(1.15), int(width * .22))
    column_width = (width - first_width) / len(categories)
    font_size = 8 if len(categories) > 10 or len(series) > 5 else 9
    # Reserve actual text lines, including Chinese station names, within the original box.
    chars_per_line = max(1, int(column_width / Pt(font_size)))
    header_lines = max(math.ceil(len(name) / chars_per_line) for name in categories)
    row_lines = [max(1, math.ceil(len(s.name or '') / max(1, int(first_width / Pt(font_size))))) for s in series]
    line_height = Pt(font_size * 1.25)
    header_height = int(header_lines * line_height + Pt(6))
    row_heights = [int(lines * line_height + Pt(5)) for lines in row_lines]
    table_height = header_height + sum(row_heights)
    if table_height > height * .65:
        _add_overflow_tables(prs, slide, shape, categories, series)
        return
    shape.height = height - table_height - Pt(5)
    table_shape = slide.shapes.add_table(len(series) + 1, len(categories) + 1,
        shape.left, shape.top + shape.height + Pt(5), width, table_height)
    table_shape.name = DATA_TABLE_PREFIX + str(shape.shape_id)
    table = table_shape.table
    table.columns[0].width = first_width
    remaining = width - first_width
    for index in range(len(categories)):
        table.columns[index + 1].width = int(remaining / len(categories))
    table.columns[len(categories)].width += width - sum(c.width for c in table.columns)
    rows = [['指标'] + categories] + [[s.name or '问题数量'] + [metric_text(v) for v in s.values] for s in series]
    for row_index, values in enumerate(rows):
        table.rows[row_index].height = header_height if row_index == 0 else row_heights[row_index - 1]
        for col_index, value in enumerate(values):
            cell = table.cell(row_index, col_index)
            cell.text = value
            cell.margin_left = cell.margin_right = Pt(2)
            cell.margin_top = cell.margin_bottom = Pt(1)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string('EDF2F7' if row_index == 0 else 'FFFFFF')
            for paragraph in cell.text_frame.paragraphs:
                if row_index and col_index == 0:
                    paragraph.clear()
                    key = paragraph.add_run(); key.text = '■ '
                    try:
                        key.font.color.rgb = series[row_index - 1].format.fill.fore_color.rgb
                    except (ValueError, AttributeError, TypeError):
                        key.font.color.rgb = RGBColor.from_string('64748B')
                    label = paragraph.add_run(); label.text = value
                paragraph.alignment = PP_ALIGN.CENTER
                paragraph.space_before = paragraph.space_after = Pt(0)
                paragraph.line_spacing = Pt(font_size * 1.25)
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)
                    run.font.name = FONT
                    run.font.bold = row_index == 0
                    if run.text != '■ ':
                        run.font.color.rgb = RGBColor.from_string('334155')
                    _font_properties(run._r.get_or_add_rPr())
            for edge in ('lnL','lnR','lnT','lnB'):
                line = OxmlElement('a:' + edge); line.set('w',str(Pt(.5)))
                fill = OxmlElement('a:solidFill'); color = OxmlElement('a:srgbClr'); color.set('val','CBD5E1')
                fill.append(color); line.append(fill); cell._tc.get_or_add_tcPr().append(line)
    for node in chart._chartSpace.xpath('.//c:dTable'):
        node.getparent().remove(node)
    chart.has_legend = False
    # Category names remain visible in the real table even if an old chart engine skips zeroes.
    if chart.chart_type in {XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.COLUMN_STACKED}:
        chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.NONE


def normalize_presentation(prs):
    for slide in list(prs.slides):
        for shape in list(slide.shapes):
            if getattr(shape,'has_chart',False):
                chart = shape.chart
                add_editable_value_table(prs, slide, shape)
                for node in chart._chartSpace.xpath('.//c:axId | .//c:crossAx'):
                    node.set('val', str(int(node.get('val')) % (2 ** 32)))
                for node in chart._chartSpace.xpath('.//c:numCache/c:formatCode'):
                    node.text = '0.########'
                for node in chart._chartSpace.xpath('.//c:externalData/c:autoUpdate'):
                    node.set('val','0')
                if chart.chart_type in BAR_TYPES:
                    chart.value_axis.minimum_scale = 0
                    if not any(float(v or 0) for s in chart.series for v in s.values):
                        chart.value_axis.maximum_scale = 1
                    chart.value_axis.tick_labels.number_format = '0.########'
                    chart.value_axis.tick_labels.number_format_is_linked = False
                for plot in chart.plots:
                    if plot.has_data_labels:
                        labels = plot.data_labels
                        labels.number_format_is_linked = False
                        if chart.chart_type == XL_CHART_TYPE.PIE:
                            labels.show_value = False
                            labels.show_category_name = False
                            labels.show_series_name = False
                            labels.show_percentage = True
                            labels.number_format = '0%'
                        else:
                            labels.show_value = True
                            labels.number_format = '0.########'
                for node in chart._chartSpace.xpath('.//a:rPr | .//a:defRPr | .//a:endParaRPr'):
                    _font_properties(node)
            for node in shape._element.xpath('.//a:rPr | .//a:defRPr | .//a:endParaRPr'):
                _font_properties(node)
    prs.core_properties.version = COMPATIBILITY_VERSION
