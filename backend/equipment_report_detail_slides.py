"""Native, editable equipment detail pages following the supplied template layouts."""
from copy import deepcopy
import colorsys
import math
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from pptx.opc.packuri import PackURI
from non_oil_report_presentation import _delete_slide, _move_slide, _remove_shape, _add_picture_contain, _set_chart_fonts
from equipment_report_analysis import canonical_unit, unit_order


def wrap(text, width, size):
    capacity = max(4, (width*72-8)/size)
    result = []
    for paragraph in str(text).split('\n'):
        line, used = '', 0
        for char in paragraph:
            units = .55 if ord(char) < 128 else 1
            if used+units > capacity:
                result.append(line)
                line, used = '', 0
            line += char
            used += units
        result.append(line)
    return result


def pages(text, width, height, size):
    lines = wrap(text, width, size)
    per = max(1, math.floor(height*72/(size*1.35)))
    return ['\n'.join(lines[i:i+per]) for i in range(0, len(lines), per)] or ['']


def text(slide, value, x, y, w, h, size=18, color='111111', bold=False, fill=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    frame = shape.text_frame
    frame.margin_top = frame.margin_bottom = Pt(0)
    frame.margin_left = frame.margin_right = Pt(2)
    frame.word_wrap = True
    frame.text = value
    for paragraph in frame.paragraphs:
        paragraph.font.name = 'Microsoft YaHei'
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor.from_string(color)
        paragraph.line_spacing = 1.2
        paragraph.space_after = Pt(0)
    return shape


def new_page(prs, source, suffix=None, ai=False):
    slide = prs.slides.add_slide(source.slide_layout)
    for shape in list(slide.shapes):
        _remove_shape(shape)
    for shape in source.shapes:
        if shape.has_text_frame and shape.top < Inches(.9):
            slide.shapes._spTree.insert_element_before(deepcopy(shape._element), 'p:extLst')
    if suffix:
        for shape in list(slide.shapes):
            _remove_shape(shape)
        text(slide, '三、检查发现', .69, .24, 2.5, .6, 28, bold=True)
        text(slide, '——'+suffix, 3.18, .24, 9.1, .6, 26, 'C00000', True)
    if ai:
        text(slide, 'AI辅助选题；选题证据及人工调整详见报告面板', .35, 7.18, 11.8, .2, 9, '64748B')
    _move_slide(prs, slide, len(prs.slides)-2)
    return slide


def photos(slide, issues, box, storage_root):
    x, y, w, h = box
    if not issues:
        text(slide, '暂无对应选题', x, y, w, h, 18, '64748B')
        return
    width = (w-.12*(len(issues)-1))/len(issues)
    for i, issue in enumerate(issues):
        first = len(slide.shapes)
        _add_picture_contain(slide, issue.get('issue_photo'),
            tuple(Inches(v) for v in (x+i*(width+.12), y, width, h-.35)), storage_root)
        for shape in list(slide.shapes)[first:]:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = RGBColor.from_string('64748B')
        text(slide, f"#{issue['issue_id']} {issue.get('station_name', '')}",
             x+i*(width+.12), y+h-.32, width, .3, 11)


def build_details(prs, original, report, storage_root):
    analysis = report.get('equipment_analysis') or {}
    issues = analysis.get('issues') or []
    distribution = analysis.get('phrase_distribution') or []
    for old in original[10:39]:
        _delete_slide(prs, old)
    # python-pptx allocates new slide part names from the current slide count.
    # Compact survivors first: the retained slide40 would otherwise collide once
    # enough evidence pages are added. Relationships follow the part objects.
    for index, existing in enumerate(prs.slides, 1):
        existing.part.partname = PackURI(f'/ppt/slides/slide{index}.xml')
    # Keep every phrase: dense legends continue rather than hiding categories under "other".
    slide = new_page(prs, original[10])
    text(slide, '加油站设备设施各类问题占比情况', .4, 1.03, 12.4, .5, 22, '0000FF', True)
    data = CategoryChartData()
    data.categories = [r['name'] for r in distribution] or ['暂无问题']
    data.add_series('问题数', [r['count'] for r in distribution] or [0])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.PIE, Inches(.35), Inches(1.7), Inches(8.6), Inches(5.25), data).chart
    chart.has_legend = False
    chart.has_title = False
    plot = chart.plots[0]
    colors = [''.join(f'{round(c*255):02X}' for c in colorsys.hls_to_rgb((i*.618034) % 1, .48, .65))
              for i in range(len(distribution))]
    for point, color in zip(chart.series[0].points, colors):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = RGBColor.from_string(color)
    plot.has_data_labels = True
    plot.data_labels.show_percentage = True
    plot.data_labels.show_value = False
    plot.data_labels.show_category_name = len(distribution) <= 10
    plot.data_labels.position = XL_LABEL_POSITION.BEST_FIT
    if len(distribution) > 15:
        plot.has_data_labels = False
    _set_chart_fonts(chart, category_size=11)
    maximum = distribution[0]['count'] if distribution else 0
    leaders = [r for r in distribution if r['count'] == maximum]
    if leaders:
        lead = leaders[0]
        note = f"{lead['name']}是单类问题{'并列' if len(leaders)>1 else ''}最多的项，共 {lead['count']} 个问题，占比 {lead['percentage']}%，是本次巡检最集中的单一问题。"
    else:
        note = '本次报告问题库暂无参与统计的问题。'
    note_pages = pages(note, 3.65, 2.15, 17)
    text(slide, note_pages[0], 9.25, 1.7, 3.65, 2.15, 17, 'C00000', True)
    legend_slide, x, y, bottom, width = slide, 9.25, 4.0, 7.05, 3.65
    for row, color in zip(distribution, colors):
        label = '\n'.join(wrap(f"{row['name']}：{row['count']}项 / {row['percentage']}%", width-.3, 12))
        height = max(.32, len(label.splitlines())*12*1.35/72+.08)
        if y+height > bottom:
            legend_slide = new_page(prs, original[10])
            text(legend_slide, '检查内容短语统计（图例续页）', .5, 1.05, 12, .5, 22, '0000FF', True)
            x, y, width = .7, 1.8, 11.8
            label = '\n'.join(wrap(f"{row['name']}：{row['count']}项 / {row['percentage']}%", width-.3, 12))
            height = max(.32, len(label.splitlines())*12*1.35/72+.08)
        text(legend_slide, '', x, y+.04, .18, .18, fill=color)
        text(legend_slide, label, x+.28, y, width-.3, height, 12)
        y += height
    for value in note_pages[1:]:
        extra = new_page(prs, original[10])
        text(extra, '检查内容短语统计（续）', .5, 1.05, 12, .5, 22, '0000FF', True)
        text(extra, value, .7, 1.8, 11.8, 5.2, 18)

    groups = analysis.get('high_groups') or []
    for start in range(0, max(1, len(groups)), 2):
        pair = groups[start:start+2]
        causes = [pages('原因（参考表原文）：'+ '；'.join(g['causes']), 5.95, 2.0, 17) for g in pair]
        titles = [pages(f"{start+i+1}、{g['phrase']}（{len(g['issues'])}项）", 5.95, .72, 20) for i,g in enumerate(pair)]
        for continuation in range(max([len(p) for p in causes+titles] or [1])):
            slide = new_page(prs, original[11], ai=True)
            text(slide, '加油站设备设施高频问题原因分析', .49, 1.0, 12, .55, 23, '0000FF', True)
            if not pair:
                text(slide, '本次暂无经AI确认的高频问题。', .8, 2, 11, 1)
            for i, group in enumerate(pair):
                x = 3.69 if len(pair) == 1 else .35+i*6.65
                text(slide, titles[i][min(continuation,len(titles[i])-1)], x, 1.7, 5.95, .72, 20, bold=True)
                representatives = sorted(group['issues'], key=lambda r: (not bool(r.get('issue_photo')), r['issue_id']))[:2]
                photos(slide, representatives, (x, 2.5, 5.95, 2.35), storage_root)
                text(slide, causes[i][continuation] if continuation<len(causes[i]) else '', x, 5.02, 5.95, 2, 17)

    special_ids = set(analysis.get('special_issue_ids') or [])
    for region in sorted(report.get('region_rows') or [], key=lambda r: unit_order(r['unit_name'])):
        unit = canonical_unit(region['unit_name'])
        stations = [s for s in report.get('station_ranking', []) if canonical_unit(s['management_unit']) == unit]
        unit_issues = [i for i in issues if i['management_unit'] == unit]
        special = [i for i in unit_issues if i['issue_id'] in special_ids]
        stats = f"检查站数：{region['station_count']}站\n问题总数：{region['issue_count']}项\n站均问题数：{region['average_issue_count']}项\n"+'、'.join(s['station_name'] for s in stations)
        stat_pages = pages(stats, 6.1, 2.1, 18)
        # Two evidence cards per page; all manually selected special issues are retained.
        chunks = [special[i:i+2] for i in range(0, len(special), 2)] or [[]]
        for page_index in range(max(len(stat_pages), len(chunks))):
            chosen = chunks[page_index] if page_index<len(chunks) else []
            descriptions = pages('特性问题：\n'+'\n'.join(f"#{i['issue_id']} {i['station_name']}：{i['description']}" for i in chosen), 6.1, 2.1, 16)
            for detail in descriptions:
                slide = new_page(prs, original[13], unit, ai=True)
                text(slide, stat_pages[min(page_index,len(stat_pages)-1)], .3, 1.2, 6.1, 2.1, 18, bold=True)
                text(slide, detail if chosen else '本片区暂无选定的特性问题。', 6.9, 1.2, 6.1, 2.1, 16, fill='E8EBF2')
                text(slide, '主要问题类型：特性问题照片', .4, 3.6, 12, .5, 22, '0000FF', True)
                photos(slide, chosen, (.45, 4.2, 12.4, 2.85), storage_root)
        for station in stations:
            station_issues = [i for i in unit_issues if i.get('station_id') == station['station_id']]
            station_special = [i for i in station_issues if i['issue_id'] in special_ids]
            description = '主要问题：\n'+'\n'.join(f"#{i['issue_id']} {i['description']}" for i in station_issues) if station_issues else '本次未发现参与报告的问题。'
            main_pages = pages(description, 4.5, 4.5, 17)
            photo_chunks = [station_special[i:i+2] for i in range(0,len(station_special),2)] or [[]]
            for page_index in range(max(len(main_pages),len(photo_chunks))):
                slide = new_page(prs, original[14], unit, ai=True)
                text(slide, '', .15, 1.15, 4.8, 5.9, fill='E8EBF2')
                text(slide, station['station_name'], .3, 1.35, 4.5, .65, 24, '0000FF', True)
                text(slide, f"问题总数：{len(station_issues)}个", .3, 2.08, 4.5, .45, 20, bold=True)
                text(slide, main_pages[page_index] if page_index<len(main_pages) else '主要问题详见本站前页。', .3, 2.65, 4.5, 4.4, 17)
                text(slide, '特性问题', 5.3, 1.25, 7.5, .6, 24, bold=True)
                selected = photo_chunks[page_index] if page_index<len(photo_chunks) else []
                for i, item in enumerate(selected):
                    text(slide, f"#{item['issue_id']} {item['phrase']}", 5.25+i*3.9, 1.95, 3.7, 1, 17)
                photos(slide, selected, (5.25, 3.05, 7.6, 3.95), storage_root)

    severe = [i for i in issues if i['issue_id'] in set(analysis.get('severe_issue_ids') or [])]
    for start in range(0, max(1,len(severe)), 3):
        selected = severe[start:start+3]
        slide = new_page(prs, original[38], ai=True)
        if not selected:
            text(slide, '本次暂无选定的严重问题。', .8, 2, 11, 1)
        width = (12.6-.3*max(0,len(selected)-1))/max(1,len(selected))
        for i, issue in enumerate(selected):
            x = .35+i*(width+.3)
            # Full descriptions remain in station pages and the evidence panel.
            text(slide, '\n'.join(wrap(issue['phrase'], width, 18)), x, 1.25, width, 1.15, 18, bold=True)
            photos(slide, [issue], (x, 2.5, width, 4.0), storage_root)
            text(slide, issue['management_unit'], x, 6.6, width, .45, 14)
