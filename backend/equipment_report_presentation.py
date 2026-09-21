"""Editable equipment template with data-driven summaries and evidence detail pages."""
from copy import deepcopy
import math
from pathlib import Path
import re

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from non_oil_report_presentation import _fill_table, _remove_shape, _move_slide, _set_chart_fonts
from pptx_compatibility import _normalize_package_fonts

TEMPLATE_FILE = Path(__file__).parent / 'assets/equipment_report_template/template.pptx'
RENDERER_VERSION = 'equipment-native-7'


def normalize_template_fonts(prs):
    for part in prs.part.package.iter_parts():
        root = getattr(part, '_element', None)
        if root is None:
            continue
        for font in root.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}buFont'):
            bullet = font.getparent().find('{http://schemas.openxmlformats.org/drawingml/2006/main}buChar')
            if font.get('typeface') == 'Wingdings' and bullet is not None and bullet.get('char') == 'l':
                bullet.set('char', '\u25cf')
            for attribute in ('charset', 'panose', 'pitchFamily'):
                font.attrib.pop(attribute, None)
    _normalize_package_fonts(prs)


def replace_text(shape, replacements):
    # Match across rich-text runs, retaining the formatting outside replaced spans.
    for paragraph in shape.text_frame.paragraphs:
        runs = list(paragraph.runs)
        for pattern, replacement in replacements:
            text = ''.join(run.text for run in runs)
            for match in reversed(list(re.finditer(pattern, text))):
                start, end = match.span()
                positions, cursor = [], 0
                for run in runs:
                    positions.append((cursor, cursor + len(run.text)))
                    cursor += len(run.text)
                inserted = False
                for run, (left, right) in zip(runs, positions):
                    if right <= start or left >= end:
                        continue
                    prefix = run.text[:max(0, start-left)]
                    suffix = run.text[max(0, end-left):] if right > end else ''
                    run.text = prefix + (replacement if not inserted else '') + suffix
                    inserted = True


def update_chart(slide, rows, keys):
    shape = next(shape for shape in slide.shapes if shape.has_chart)
    chart = shape.chart
    data = CategoryChartData()
    data.categories = [str(row['name']) for row in rows] or ['暂无数据']
    for name, key in keys:
        data.add_series(name, [float(row.get(key) or 0) for row in rows] or [0])
    # Template charts reference the author's local workbook. Embed fresh data instead.
    for node in chart._chartSpace.xpath('./c:externalData'):
        relation = node.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        node.getparent().remove(node)
        if relation:
            chart.part.drop_rel(relation)
    chart.replace_data(data)
    for node in chart._chartSpace.xpath('.//c:dLbl | .//c:dPt | .//c:extLst'):
        node.getparent().remove(node)
    chart.value_axis.minimum_scale = 0
    chart.value_axis.maximum_scale = None if any(row.get(key) for row in rows for _, key in keys) else 1
    for plot in chart.plots:
        plot.has_data_labels = True
        plot.data_labels.show_value = True
        plot.data_labels.number_format = '0.########'
        plot.data_labels.number_format_is_linked = False
    chart.category_axis.tick_labels.number_format_is_linked = False
    _set_chart_fonts(chart, category_size=9 if len(rows) > 12 else 11)
    # The source charts extend below the canvas; keep all labels above the footer.
    shape.height = min(shape.height, Inches(7.12) - shape.top)


def fill_station_page(slide, stations):
    tables = sorted([s for s in slide.shapes if s.has_table], key=lambda s: s.left)
    count = min(5, max(1, math.ceil(len(stations)/11)))
    per_column = max(1, math.ceil(len(stations)/count))
    gap = .16
    width = min(5.4, (12.8-gap*(count-1))/count)
    left = (13.333 - count*width-gap*(count-1))/2
    for i, shape in enumerate(tables):
        if i >= count:
            _remove_shape(shape)
            continue
        chunk = stations[i*per_column:(i+1)*per_column]
        rows = [[s.get('station_name') or '未命名站点', str(s.get('issue_count', 0))] for s in chunk]
        rows = rows or [['暂无已完成巡检', '0']]
        shape.left, shape.top, shape.width = Inches(left+i*(width+gap)), Inches(1.8), Inches(width)
        shape.height = Inches(min(5.25, .44*(len(rows)+1)))
        target_width = Inches(width)
        shape.table.columns[0].width = int(target_width*.73)
        shape.table.columns[1].width = target_width-shape.table.columns[0].width
        longest = max(len(str(row[0])) for row in rows)
        font_size = max(9, min(14, width*.73*72/max(longest, 1)))
        _fill_table(shape, ['站点', '问题数'], rows, font_size=font_size)
        target_height = shape.height
        shape.table.rows[0].height = Inches(.44)
        for row in list(shape.table.rows)[1:]:
            row.height = int((target_height-Inches(.44))/len(rows))
        for r, row in enumerate(shape.table.rows):
            for cell in row.cells:
                cell.margin_left = cell.margin_right = Inches(.02)
                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = 'Microsoft YaHei'
                        run.font.color.rgb = RGBColor.from_string('FFFFFF' if r == 0 else '000000')


def build_equipment_template_presentation(report, output_path, storage_root=None):
    prs = Presentation(TEMPLATE_FILE)
    original = list(prs.slides)
    if len(prs.slides) != 40:
        raise ValueError('设备设施报告模板页数异常。')
    summary = report.get('summary') or {}
    total = int(summary.get('total_issue_count') or 0)
    stations = int(summary.get('station_count') or 0)
    average = f'{total/stations:.1f}' if stations else '0.0'
    month = str(report.get('month') or '')
    if not re.fullmatch(r'\d{4}-\d{2}', month):
        raise ValueError('设备设施报告缺少有效报告月份。')
    for shape in prs.slides[0].shapes:
        if shape.has_text_frame:
            replace_text(shape, [(r'\d{4}年\s*\d{1,2}月', f'{month[:4]}年{int(month[5:])}月')])
    for shape in prs.slides[3].shapes:
        if shape.has_text_frame:
            replace_text(shape, [(r'199', str(total)), (r'3\.7', average)])

    regions = report.get('region_rows') or []
    units = sum(row.get('unit_type') == 'holding' for row in regions)
    intro = next(s for s in prs.slides[5].shapes if s.has_text_frame and '本月完成' in s.text)
    replace_text(intro, [(r'7(?=个片区)', str(len(regions)-units)), (r'4(?=家股权单位)', str(units)),
                        (r'54(?=座加油站)', str(stations)), (r'3\.7', average), (r'199', str(total))])
    intro.height = Inches(.76)
    for paragraph in intro.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(19)
    chart_shape = next(s for s in prs.slides[5].shapes if s.has_chart)
    chart_shape.top = Inches(1.8)
    chart_shape.height = Inches(5.25)
    update_chart(prs.slides[5], [dict(row, name=row['unit_name']) for row in regions],
                 [('问题总数', 'issue_count'), ('受检加油站数量', 'station_count'), ('单站平均问题数', 'average_issue_count')])
    update_chart(prs.slides[8], report.get('area_distribution') or [], [('问题数', 'count'), ('占比', 'percentage')])
    update_chart(prs.slides[9], report.get('item_distribution') or [], [('问题数', 'count'), ('占比', 'percentage')])
    ranking = sorted(report.get('station_ranking') or [], key=lambda s: (-int(s.get('issue_count') or 0), s.get('station_name') or ''))
    source = prs.slides[6]
    # Never omit stations: extremely large selections continue on another ranking page.
    for offset in range(55, len(ranking), 55):
        page = prs.slides.add_slide(source.slide_layout)
        for shape in list(page.shapes):
            _remove_shape(shape)
        for shape in source.shapes:
            page.shapes._spTree.insert_element_before(deepcopy(shape._element), 'p:extLst')
        fill_station_page(page, ranking[offset:offset+55])
        _move_slide(prs, page, 6+offset//55)
    fill_station_page(source, ranking[:55])
    if 'equipment_analysis' in report:
        from equipment_report_detail_slides import build_details
        build_details(prs, original, report, storage_root)
    for page in prs.slides:
        for shape in page.shapes:
            if shape.has_text_frame:
                replace_text(shape, [(r'严重问题', '重点问题')])
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    # Keep the existing font whitelist without adding compatibility pages or moving template objects.
    normalize_template_fonts(prs)
    prs.save(output_path)
    return {'ppt_path': str(output_path), 'slide_count': len(prs.slides)}
