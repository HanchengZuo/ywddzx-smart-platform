"""Update native template elements without replacing their layouts or chart types."""
from contextlib import contextmanager
from copy import deepcopy
import math
from lxml import etree
from pptx.oxml import parse_xml
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from pptx_compatibility import _font_properties

FONT = 'Microsoft YaHei'


def paragraph(shape, spans, size, *, first=False, bold=False, props=None, align=None, bullet=False):
    p = shape.text_frame.paragraphs[0] if first else shape.text_frame.add_paragraph()
    p.clear()
    old = p._p.find(qn('a:pPr'))
    if old is not None: p._p.remove(old)
    if props is not None: p._p.insert(0, deepcopy(props))
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    pr = p._p.get_or_add_pPr()
    if bullet:
        for node in list(pr):
            if node.tag in {qn('a:buNone'), qn('a:buFont'), qn('a:buChar')}:
                pr.remove(node)
        font = OxmlElement('a:buFont'); font.set('typeface', FONT)
        char = OxmlElement('a:buChar'); char.set('char', '■' if bullet=='square' else '▶')
        # OOXML requires bullet properties before tabs/default run properties.
        before = next((n for n in pr if n.tag in {qn('a:tabLst'), qn('a:defRPr'), qn('a:extLst')}), None)
        index = pr.index(before) if before is not None else len(pr)
        pr.insert(index, font); pr.insert(index+1, char)
        pr.set('marL', str(Pt(size * 1.1))); pr.set('indent', str(-Pt(size * 1.1)))
    if align is not None: p.alignment = align
    for text, color in spans:
        run = p.add_run(); run.text = text
        run.font.name = FONT; run.font.size = Pt(size); run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
        _font_properties(run._r.get_or_add_rPr())
    return p


def fit_size(texts, width, height, maximum=18, minimum=11, leading=1.45, reserved=0):
    """Keep the complete paragraph above the next template element."""
    from unicodedata import east_asian_width
    for size in range(maximum, minimum-1, -1):
        columns = max(1, (width * 72 - 15) / size)
        lines = sum(max(1, math.ceil(sum(1 if east_asian_width(c) in 'WF' else .55 for c in text) / columns)) for text in texts)
        if lines * size * leading + reserved <= height * 72: return size
    return minimum


@contextmanager
def editable_plot(chart):
    """python-pptx cannot address pie3DChart; temporarily use its 2D XML adapter.

    The saved chart remains the original editable 3D chart, including view3D,
    lighting and depth. This does not flatten it into an image.
    """
    plots = []
    for original in chart._chartSpace.xpath('.//c:pie3DChart'):
        clone = deepcopy(original); clone.tag = qn('c:pieChart')
        # Reparse to obtain python-pptx's CT_PieChart rather than a generic
        # lxml element, whose XPath does not know OOXML namespace prefixes.
        plot = parse_xml(etree.tostring(clone))
        original.getparent().replace(original, plot)
        plots.append(plot)
    try: yield chart
    finally:
        for plot in plots: plot.tag = qn('c:pie3DChart')


def update_chart(shape, rows, series, *, pie=False, title=None):
    data = CategoryChartData(number_format='0')
    data.categories = [row.get('unit_name') or row.get('name') or '' for row in rows] or ['暂无数据']
    for label, key in series:
        data.add_series(label, [int(row.get(key) or 0) for row in rows] or [0], number_format='0')
    with editable_plot(shape.chart) as chart:
        chart.replace_data(data)
        # Old per-point labels/legend overrides can carry template values or
        # category indexes that are no longer valid after replacing the data.
        for node in chart._chartSpace.xpath('.//c:dLbl | .//c:ser/c:dLbls | .//c:legendEntry'):
            node.getparent().remove(node)
        if title:
            chart.has_title = True
            chart.chart_title.has_text_frame = True
            chart.chart_title.text_frame.text = title
            for p in chart.chart_title.text_frame.paragraphs:
                for r in p.runs:
                    r.font.name = FONT; r.font.size = Pt(14); r.font.bold = True
                    r.font.color.rgb = RGBColor.from_string('595959')
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.TOP if pie else XL_LEGEND_POSITION.BOTTOM
        chart.legend.font.name = FONT; chart.legend.font.size = Pt(11 if pie else 9)
        for plot in chart.plots:
            plot.has_data_labels = bool(rows)
            labels = plot.data_labels
            labels.show_value = not pie; labels.show_percentage = pie
            labels.show_category_name = False; labels.show_series_name = False
            labels.number_format = '0.00%' if pie else '0'
            labels.number_format_is_linked = False
            labels.font.name = FONT; labels.font.size = Pt(11 if pie else 9)
            labels.position = XL_LABEL_POSITION.OUTSIDE_END
            for body in chart._chartSpace.xpath('.//c:dLbls/c:txPr/a:bodyPr'):
                body.set('wrap', 'none'); body.set('horzOverflow', 'overflow')
        if pie:
            colors = ['4F81BD', 'C0504D', '9BBB59', '8064A2', '4BACC6', 'F79646', 'A5A5A5', '264478']
            for i, point in enumerate(chart.series[0].points):
                point.format.fill.solid(); point.format.fill.fore_color.rgb = RGBColor.from_string(colors[i % len(colors)])
            for layout in chart._chartSpace.xpath('.//c:legend/c:layout | .//c:plotArea/c:layout'):
                for node in list(layout): layout.remove(node)
            layout=chart._chartSpace.xpath('.//c:plotArea/c:layout')[0]
            manual=OxmlElement('c:manualLayout')
            for tag, value in [('xMode','factor'),('yMode','factor'),('wMode','factor'),('hMode','factor'),
                ('x','.12'),('y','.12'),('w','.74'),('h','.76')]:
                node=OxmlElement('c:'+tag);node.set('val',value);manual.append(node)
            layout.append(manual)
        else:
            for i, color in enumerate(['4F81BD', 'C0504D']):
                chart.series[i].format.fill.solid()
                chart.series[i].format.fill.fore_color.rgb = RGBColor.from_string(color)
            chart.value_axis.minimum_scale = 0
            if not any(int(row.get(key) or 0) for row in rows for _, key in series):
                chart.value_axis.maximum_scale = 1
            chart.value_axis.tick_labels.number_format = '0'
            chart.value_axis.tick_labels.number_format_is_linked = False
            chart.value_axis.major_gridlines.format.line.color.rgb = RGBColor.from_string('D9D9D9')
            chart.value_axis.major_gridlines.format.line.width = Pt(.5)
            chart.category_axis.tick_labels.font.name = FONT
            chart.category_axis.tick_labels.font.size = Pt(9)
        for node in chart._chartSpace.xpath('.//c:numCache/c:formatCode'):
            node.text = '0'
        for node in chart._chartSpace.xpath('.//c:axId | .//c:crossAx'):
            node.set('val', str(int(node.get('val')) % (2 ** 32)))


def frame_props(shape, index=0):
    return deepcopy(shape.text_frame.paragraphs[index]._p.find(qn('a:pPr')))


def split_arrow(slide, x, y, size=.18):
    """Wingdings' half-hollow arrow as native geometry, not a font fallback."""
    for suffix, points, color in [('upper',[(0,0),(100,50),(22,50)],'FFFFFF'),
            ('lower',[(0,100),(100,50),(22,50)],'000000')]:
        builder=slide.shapes.build_freeform(0,0,scale=Inches(size)/100)
        builder.move_to(*points[0]);builder.add_line_segments(points[1:],close=True)
        shape=builder.convert_to_shape(Inches(x),Inches(y));shape.name='safety-split-arrow-'+suffix
        shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(color)
        shape.line.color.rgb=RGBColor.from_string('000000');shape.line.width=Pt(.65)


def clone_pie_chart(slide, source, x, y, w, h):
    # Each chart owns its workbook; sharing a ChartPart would silently replace
    # the distribution page's data when updating the analysis page.
    data=CategoryChartData();data.categories=['暂无数据'];data.add_series('问题数量',[0])
    shape=slide.shapes.add_chart(XL_CHART_TYPE.PIE,Inches(x),Inches(y),Inches(w),Inches(h),data)
    chart=shape.chart
    external=deepcopy(chart._chartSpace.find(qn('c:externalData')))
    for child in list(chart._chartSpace):chart._chartSpace.remove(child)
    for child in source.chart._chartSpace:
        if child.tag!=qn('c:externalData'):chart._chartSpace.append(deepcopy(child))
    chart._chartSpace.append(external)
    return shape


def compact_pie_layout(slide, shape, rows, title):
    """Keep the title, editable 3D plot and legend in explicit, adjacent bands.

    Office/WPS/LibreOffice reserve different amounts of whitespace for a native
    3D chart's title and legend. Independent native text/shapes avoid that gap.
    """
    shape.chart.has_title=False;shape.chart.has_legend=False
    for node in shape.chart._chartSpace.xpath('.//c:plotArea/c:layout'):
        for child in list(node):node.remove(child)
    x,w=shape.left/914400,shape.width/914400
    heading=slide.shapes.add_textbox(Inches(x),Inches(1.36),Inches(w),Inches(.42))
    heading.name='safety-pie-title'
    paragraph(heading,[(title,'595959')],16,first=True,bold=True,align=PP_ALIGN.CENTER)
    columns=3 if len(rows)>6 else 2
    legend_rows=max(1,math.ceil(len(rows)/columns))
    from unicodedata import east_asian_width
    cell=w/columns
    row_heights=[]
    font_size=9 if len(rows)>6 else 10
    for row in range(legend_rows):
        names=rows[row*columns:(row+1)*columns]
        lines=max([math.ceil(sum(1 if east_asian_width(c) in 'WF' else .55 for c in item['name']) * font_size / ((cell-.15)*72)) for item in names] or [1])
        row_heights.append(max(.24,lines*font_size*1.25/72+.03))
    legend_height=sum(row_heights)
    legend_top=7.12-legend_height
    shape.top=Inches(1.86);shape.height=Inches(legend_top-1.94)
    palette=['4F81BD','C0504D','9BBB59','8064A2','4BACC6','F79646','A5A5A5','264478']
    for i,row in enumerate(rows):
        left=x+(i%columns)*cell;top=legend_top+sum(row_heights[:i//columns])
        key=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(left),Inches(top+.07),Inches(.09),Inches(.09))
        key.name='safety-pie-key';key.fill.solid();key.fill.fore_color.rgb=RGBColor.from_string(palette[i%len(palette)]);key.line.fill.background()
        label=slide.shapes.add_textbox(Inches(left+.13),Inches(top),Inches(cell-.15),Inches(row_heights[i//columns]))
        label.name='safety-pie-label'
        label.text_frame.margin_left=label.text_frame.margin_right=0
        label.text_frame.margin_top=label.text_frame.margin_bottom=0
        paragraph(label,[(row['name'],'595959')],font_size,first=True)
    if len(rows)>10:
        labels=shape.chart._chartSpace.xpath('.//c:pie3DChart/c:dLbls')[0]
        for index,row in enumerate(rows):
            if row.get('percentage',100)>=3:continue
            entry=OxmlElement('c:dLbl');idx=OxmlElement('c:idx');idx.set('val',str(index));entry.append(idx)
            hidden=OxmlElement('c:delete');hidden.set('val','1');entry.append(hidden);labels.insert(0,entry)
