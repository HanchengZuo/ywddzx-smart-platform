from datetime import date
import unittest
from unittest.mock import patch, MagicMock
from flask import Flask
from safety_video_themes import VIDEO_POLICY, VIDEO_GUIDANCE, normalize_video_groups, selected_video_group, selected_video_issue
from safety_report_analysis import analysis_context, analyze, build_report, normalize_settings, validate_analysis
from safety_report_data import UNIT_ORDER, REGION_ORDER, unit_key, unit_label, serialize, scope
from safety_report_http import register_safety_report
from report_ai_memory import fingerprint
from tests.test_safety_report import raw_issue, mock_analysis


class UnitOrderingTest(unittest.TestCase):
    def test_all_declared_units_keep_fixed_order_without_counting_equity_as_regions(self):
        expected=['浦东','闵普徐','松金','嘉青','南汇','宝静','奉贤','崇明','中油奉贤','中油同盛',
            '中油康桥','中油农工商','中油上海','中油港汇','中石油上港','中油浦东','中油华鑫','中油中燃（海滨油库）']
        self.assertEqual(UNIT_ORDER,expected)
        self.assertEqual(sorted(list(reversed(expected)),key=unit_key),expected)
        self.assertEqual(unit_key('中油中燃公司'),unit_key('中油中燃（海滨油库）'))
        self.assertEqual(unit_label('浦东管理片区'),'浦东片区')
        rows=[serialize(raw_issue(i,region=name+'公司',station_id=i)) for i,name in enumerate(expected,1)]
        result=scope(rows,[],'onsite','2026年09月')
        self.assertEqual([r['unit_name'] for r in result['units']],expected)
        self.assertEqual(result['region_count'],8)
        self.assertEqual(result['holding_unit_count'],10)
        subset=scope([rows[15],rows[8],rows[0]],[],'onsite','2026年09月')
        self.assertEqual([r['unit_name'] for r in subset['units']],['浦东','中油奉贤','中油浦东'])
        self.assertEqual(len(REGION_ORDER),8)

    def test_unlisted_real_units_are_not_lost_or_guessed_from_station_names(self):
        rows=[serialize(raw_issue(1,region='中油浦江公司',station_name='浦东站')),
            serialize(raw_issue(2,region='中油奉贤公司'))]
        result=scope(rows,[],'onsite','2026年09月')
        self.assertEqual([r['unit_name'] for r in result['units']],['中油奉贤','中油浦江'])


class VideoThemeTest(unittest.TestCase):
    def test_synonyms_merge_without_combining_unrelated_risks_or_changing_ids(self):
        groups=[{'name':name,'summary':name+'的实际表现','issue_ids':[i]} for i,name in enumerate([
            '检维修未布置消防隔离','检维修防护隔离缺失','检修现场未设隔离与消防器材','安全帽未佩戴','连接管线顺序不规范'],1)]
        value=normalize_video_groups({'groups':groups,'causes':['可能检查不足'],'complete':True})
        self.assertEqual(value['groups'][0]['name'],'作业隔离与消防防护不到位')
        self.assertEqual(value['groups'][0]['issue_ids'],[1,2,3])
        self.assertEqual(len(value['groups']),3)
        validate_analysis({**value,'suggestions':[]},[{'issue_id':i} for i in range(1,6)])

    def test_only_video_typical_gets_new_context_and_coarse_guidance(self):
        row=serialize(raw_issue(1,'video'))
        context=analysis_context([row],'video','__typical__')
        self.assertEqual(context['classification_policy'],VIDEO_POLICY)
        self.assertEqual(context['grouping_guidance'],VIDEO_GUIDANCE)
        self.assertIn('不按措辞',VIDEO_GUIDANCE)
        self.assertNotIn('classification_policy',analysis_context([row],'onsite','设备设施'))
        self.assertNotIn('classification_policy',analysis_context([row],'video','卸油作业'))

    def test_synonyms_across_40_issue_batches_are_merged_and_manual_theme_drives_photos(self):
        rows=[serialize(raw_issue(i,'video')) for i in range(1,42)]
        def generate(context):
            ids=[i['issue_id'] for i in context['issues']]
            return {'payload':{'groups':[{'name':'检维修防护隔离缺失' if ids[0]==1 else '检修现场未设隔离与消防器材',
                'summary':'作业未设置隔离或消防防护','issue_ids':ids}], 'causes':['可能执行不到位'],'suggestions':[]}}
        with patch('safety_report_analysis.generate_analysis',side_effect=generate):
            result=analyze(rows,'video','__typical__')
        self.assertEqual(len(result['groups']),1)
        self.assertEqual(len(result['groups'][0]['issue_ids']),41)
        self.assertEqual(selected_video_group(rows,result)['issue_ids'],list(range(1,42)))

    def test_saved_video_theme_picks_photo_from_that_theme_and_ignores_legacy_single_issue_choice(self):
        rows=[raw_issue(1),raw_issue(2,'video',issue_photo='a.jpg'),raw_issue(3,'video',issue_photo='b.jpg')]
        def analysis(issues,mode,category,*args):
            if mode=='video':return {'groups':[{'name':'流程不规范','summary':'连接顺序不规范','issue_ids':[2]},
                {'name':'着装不规范','summary':'未按规定着装','issue_ids':[3]}],'causes':['可能培训不足']}
            return mock_analysis(issues)
        with patch('safety_report_analysis.analyze',side_effect=analysis),patch('safety_report_analysis.generate_analysis',return_value={
            'payload':{'groups':[],'causes':[],'suggestions':[{'title':'治理','content':'整改风险'}]}}):
            report=build_report(date(2026,9,1),date(2026,10,1),rows,[],{'onsite':[],'video':[]},
                normalize_settings({'typical_video_category':'着装不规范','typical_issue_ids':{'video':2}}))
        finding=next(f for f in report['typical_findings'] if f['mode']=='video')
        self.assertEqual(finding['category'],'着装不规范')
        self.assertEqual(finding['issues'][0]['issue_id'],3)

    def test_legacy_report_rerender_never_uses_photo_from_a_different_video_theme(self):
        from tests.test_safety_report import fixture_report
        from safety_report_presentation import typical_page, TEMPLATE_FILE
        from pptx import Presentation
        from pathlib import Path
        report=fixture_report(settings={'typical_issue_ids':{'video':2}})
        report['safety_issues'] += [{'issue_id':7,'mode':'video','description':'未戴帽','station_name':'测试站','issue_photo':'hat.jpg'}]
        report['safety_analysis']['video:__typical__']={'groups':[
            {'name':'作业隔离与消防防护不到位','summary':'隔离缺失','issue_ids':[2]},
            {'name':'个人劳动防护与着装不规范','summary':'未戴帽','issue_ids':[7]}]}
        report['safety_settings']['typical_video_category']='个人劳动防护与着装不规范'
        prs=Presentation(TEMPLATE_FILE)
        with patch('safety_report_presentation.photo') as photo:
            typical_page(prs.slides[5],report,Path('/tmp'))
        self.assertEqual(photo.call_args_list[-1].args[1]['issue_id'],7)

    def test_video_photo_choice_stays_in_theme_and_uses_only_photographed_issues(self):
        rows=[serialize(raw_issue(i,'video',issue_photo='' if i==3 else f'{i}.jpg')) for i in (2,1,3,4)]
        group={'issue_ids':[1,2,3]}
        for selection,expected in ((None,1),(2,2),(3,1),(4,1),(999,1)):
            self.assertEqual(selected_video_issue(rows,group,selection)['issue_id'],expected)
        self.assertIsNone(selected_video_issue(rows,{'issue_ids':[3]},2))
        self.assertIsNone(selected_video_issue(rows,None,2))

    def test_saved_manual_video_photo_reaches_ppt_without_changing_category_narrative(self):
        from tests.test_safety_report import fixture_report
        from safety_report_presentation import typical_page, TEMPLATE_FILE, video_typical_narrative
        from pptx import Presentation
        from pathlib import Path
        rows=[raw_issue(1),raw_issue(2,'video'),raw_issue(3,'video')]
        report=fixture_report('example.jpg',rows=rows,settings={'typical_issue_ids':{'video':3}})
        finding=next(f for f in report['typical_findings'] if f['mode']=='video')
        self.assertEqual(finding['issues'][0]['issue_id'],3)
        self.assertEqual(finding['issue_count'],2)
        self.assertEqual(finding['summary'],report['safety_analysis']['video:__typical__']['groups'][0]['summary'])
        self.assertIn('总计出现2次',video_typical_narrative(report))
        with patch('safety_report_presentation.photo') as photo:
            typical_page(Presentation(TEMPLATE_FILE).slides[5],report,Path('/tmp'))
        self.assertEqual(photo.call_args_list[-1].args[1]['issue_id'],3)


class VideoThemeHttpTest(unittest.TestCase):
    def test_manual_video_photo_save_validates_theme_photo_and_rolls_back_invalid_choices(self):
        conn=MagicMock(); app=Flask(__name__)
        rows=[raw_issue(1,'video',issue_photo='1.jpg'),raw_issue(2,'video',issue_photo='2.jpg'),
            raw_issue(3,'video'),raw_issue(4,'video',issue_photo='4.jpg')]
        stored={'settings':{},'revision':0}
        register_safety_report(app,connect=lambda:conn,authorize=lambda c:{'id':1},
            period=lambda *a:(date(2026,9,1),date(2026,10,1)),fetch_inputs=lambda *a:(rows,[],[]),
            has_permission=lambda *a:False,save_workspace=MagicMock(),log_usage=MagicMock())
        choice={'groups':[{'name':'隔离','summary':'防护不足','issue_ids':[1,2,3]},
            {'name':'着装','summary':'着装问题','issue_ids':[4]}],'complete':True}
        request={'settings':{'typical_issue_ids':{'video':2}},'revision':0,
            'data_signature':fingerprint([serialize(row) for row in rows])}
        def persist(cur,settings,*args):stored.update(settings=settings,revision=1)
        with patch('safety_report_http.load_settings',side_effect=lambda *args:stored.copy()),patch(
            'safety_report_http.cached_analysis',return_value=choice),patch('safety_report_http.save_settings',side_effect=persist) as save,patch(
            'safety_report_http.analyze',side_effect=AssertionError('paid AI')):
            client=app.test_client()
            response=client.put('/api/inspection-reports/safety-config',json=request)
            self.assertEqual(response.status_code,200,response.json)
            self.assertEqual(client.get('/api/inspection-reports/safety-config').json['settings']['typical_issue_ids']['video'],2)
            for invalid in (3,4,999):
                conn.reset_mock();save.reset_mock()
                request['settings']['typical_issue_ids']['video']=invalid
                response=client.put('/api/inspection-reports/safety-config',json=request)
                self.assertEqual(response.status_code,400,response.json)
                conn.rollback.assert_called_once();conn.commit.assert_not_called();save.assert_not_called()
            request['settings']['typical_issue_ids']['video']=2
            request['settings']['typical_video_category']='着装'
            self.assertEqual(client.put('/api/inspection-reports/safety-config',json=request).status_code,400)
            request['settings']['typical_issue_ids']={}
            self.assertEqual(client.put('/api/inspection-reports/safety-config',json=request).status_code,200)

    def test_only_complete_current_categories_can_be_saved_without_calling_ai(self):
        conn=MagicMock(); app=Flask(__name__); rows=[raw_issue(1,'video')]
        register_safety_report(app,connect=lambda:conn,authorize=lambda c:{'id':1},
            period=lambda *a:(date(2026,9,1),date(2026,10,1)),fetch_inputs=lambda *a:(rows,[],[]),
            has_permission=lambda *a:False,save_workspace=MagicMock(),log_usage=MagicMock())
        choice={'groups':[{'name':'着装不规范','summary':'着装问题','issue_ids':[1]}],'complete':True}
        request={'settings':{'typical_video_category':'着装不规范'},'revision':0,'data_signature':fingerprint([serialize(rows[0])])}
        with patch('safety_report_http.load_settings',return_value={'settings':{},'revision':0}),patch(
            'safety_report_http.cached_analysis',return_value=choice),patch('safety_report_http.save_settings') as save,patch('safety_report_http.analyze') as ai:
            self.assertEqual(app.test_client().put('/api/inspection-reports/safety-config',json=request).status_code,200)
            self.assertEqual(save.call_args.args[1]['typical_video_category'],'着装不规范')
            request['settings']['typical_video_category']='范围外分类'
            self.assertEqual(app.test_client().put('/api/inspection-reports/safety-config',json=request).status_code,400)
            ai.assert_not_called()
