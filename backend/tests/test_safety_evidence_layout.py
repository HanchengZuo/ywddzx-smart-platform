import copy
from io import BytesIO
from pathlib import Path
import re
import tempfile
import unittest
from unittest.mock import patch

from PIL import Image
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from safety_report_data import UNIT_ORDER, VIDEO_PAGES, unit_key, unit_label
from safety_report_presentation import (RENDERER_VERSION, TEMPLATE_FILE, backlog_page,
    build_safety_presentation, evidence_layout, evidence_page)
from safety_template_styles import evidence_text_height
from tests.test_safety_report import fixture_report, raw_issue


def geometry(shape):
    return (shape.left, shape.top, shape.width, shape.height)


def named(slide, prefix):
    return sorted((s for s in slide.shapes if s.name.startswith(prefix)), key=lambda s:s.left)


def photo_ids(slide):
    return list(map(int, re.findall(r'问题ID (\d+)', slide.notes_slide.notes_text_frame.text)))


def group(name, ids, representatives=None):
    return {'name':name, 'summary':name+'的实际共性描述。', 'issue_ids':ids,
        'representative_issue_ids':representatives or list(reversed(ids))}


class SafetyEvidenceLayoutTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.root = Path(cls.temp.name)
        for filename, size in [('wide.jpg',(640,360)), ('tall.jpg',(240,720))]:
            Image.new('RGB', size, 'steelblue').save(cls.root/filename)
        rows=[]
        for topic, ids in [('设备设施',range(1,7)), ('污染防治',range(7,10)),
                ('能力培训和意识',range(10,13)), ('风险隐患',range(13,16)),
                ('其他对生产经营有影响的问题隐患',range(16,22))]:
            for issue_id in ids:
                rows.append(raw_issue(issue_id, station_name=f'示例{issue_id}',
                    description=f'现场问题{issue_id}的实际描述。',
                    standard_detail_text=f'检查主题：{topic}\n检查内容：{topic}\n评分说明：旧评分不符合扣10分。'))
        for index,(category,_) in enumerate(VIDEO_PAGES):
            for issue_id in range(22+index*3,25+index*3):
                rows.append(raw_issue(issue_id,'video',station_name=f'示例{issue_id}加油站',
                    description=f'视频问题{issue_id}的实际描述。',
                    standard_detail_text=f'检查内容：{category}\n评分说明：旧评分不符合扣10分。'))
        cls.report=fixture_report('wide.jpg',rows=rows,settings={'focus_category':'设备设施'})
        for issue in cls.report['safety_issues']:
            issue.update(requirement=f"设施{issue['issue_id']}应保持连接可靠。",requirement_status='ready')
            if issue['issue_id']%2:issue['issue_photo']='tall.jpg'
        lookup={row['issue_id']:row for row in cls.report['safety_issues']}
        lookup[5].update(station_name='已带后缀加油站',requirement_status='manual')
        lookup[6]['requirement']='旧评分不符合扣10分。'
        lookup[6].pop('requirement_status',None)
        equipment=[group('首类',[1,2,3],[3,1,2]),group('次类',[4,5],[5,4]),group('第三类',[6])]
        for key in ('onsite:设备设施','onsite:topic:设备设施'):
            cls.report['safety_analysis'][key]['groups']=copy.deepcopy(equipment)
        for key,ids in [('onsite:污染防治',list(range(7,10))),
                ('onsite:能力培训和意识',list(range(10,13))),
                ('onsite:风险合规',list(range(13,16))),
                ('onsite:其他隐患类',list(range(16,22)))]:
            cls.report['safety_analysis'][key]['groups']=[group(key,ids)]
        for index,(category,_) in enumerate(VIDEO_PAGES):
            cls.report['safety_analysis']['video:'+category]['groups']=[
                group(category,list(range(22+index*3,25+index*3)))]
        cls.report['pending_rectification']={
            'onsite':[{'unit_name':'崇明','pending_sign':0,'pending_rectification':5},
                {'unit_name':'浦东','pending_sign':2,'pending_rectification':0}],
            'video':[{'unit_name':'松金','pending_sign':0,'pending_rectification':0},
                {'unit_name':'闵普徐','pending_sign':3,'pending_rectification':1}]}
        cls.output=cls.root/'populated-safety.pptx'
        with patch('safety_report_analysis.generate_analysis',side_effect=AssertionError('No paid AI')), patch(
                'safety_report_examples.generate_representatives',side_effect=AssertionError('No paid AI')):
            build_safety_presentation(cls.report,cls.output,cls.root)
        cls.prs=Presentation(cls.output)
        cls.template=Presentation(TEMPLATE_FILE)

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_all_31_slides_and_static_text_geometry_are_preserved(self):
        self.assertEqual(RENDERER_VERSION,'safety-native-1214-1')
        self.assertEqual(len(self.prs.slides),31)
        for index in (1,4,6,9,24,26,30):
            before=self.template.slides[index];after=self.prs.slides[index]
            self.assertEqual([(s.text,geometry(s)) for s in before.shapes if s.has_text_frame],
                [(s.text,geometry(s)) for s in after.shapes if s.has_text_frame])
        for index in range(12,24):
            before=next(s for s in self.template.slides[index].shapes if s.has_text_frame and s.top<Inches(.8))
            after=next(s for s in self.prs.slides[index].shapes if s.has_text_frame and s.top<Inches(.8))
            self.assertEqual((before.text,geometry(before)),(after.text,geometry(after)))

    def test_every_evidence_page_has_native_blue_bold_station_paragraphs(self):
        for index in range(12,24):
            with self.subTest(slide=index+1):
                captions=named(self.prs.slides[index],'safety-evidence-caption-')
                self.assertEqual(len(captions),3)
                for caption in captions:
                    paragraphs=caption.text_frame.paragraphs
                    self.assertEqual(len(paragraphs),2)
                    station=next(p for p in paragraphs if p.text.endswith('加油站'))
                    self.assertNotIn('加油站加油站',station.text)
                    self.assertEqual(station.alignment,2)
                    self.assertTrue(all(r.font.bold and str(r.font.color.rgb)=='3F95D2' for r in station.runs))
                    description=next(p for p in paragraphs if p is not station)
                    self.assertTrue(all(not r.font.bold and str(r.font.color.rgb)=='000000' for r in description.runs))
                    self.assertTrue(all(Pt(11)<=r.font.size<=Pt(16) for p in paragraphs for r in p.runs))
                    self.assertIs(station,paragraphs[0] if index>=20 else paragraphs[1])

    def test_photo_and_caption_geometry_follow_native_source_slots_without_overlap(self):
        for index in range(12,24):
            with self.subTest(slide=index+1):
                slide=self.prs.slides[index]
                _,caption_boxes,photo_boxes,norm_box,_,_,_=evidence_layout(self.template.slides[index])
                captions=named(slide,'safety-evidence-caption-')
                pictures=named(slide,'safety-evidence-photo-')
                norm=named(slide,'safety-evidence-requirements')[0]
                self.assertEqual(len(pictures),3)
                self.assertEqual((norm.left,norm.width),(norm_box[0],norm_box[2]))
                self.assertAlmostEqual(norm.top+norm.height,norm_box[1]+norm_box[3],delta=2)
                for caption,picture,cap_box,pic_box in zip(captions,pictures,caption_boxes,photo_boxes):
                    self.assertEqual((caption.left,caption.width),(cap_box[0],cap_box[2]))
                    self.assertGreater(caption.height,0)
                    self.assertLessEqual(caption.top,cap_box[1])
                    self.assertLessEqual(caption.top+caption.height+Inches(.079),norm.top)
                    self.assertEqual(picture.shape_type,13)
                    self.assertGreater(picture.height,0)
                    self.assertGreaterEqual(picture.left,pic_box[0]-1)
                    self.assertLessEqual(picture.left+picture.width,pic_box[0]+pic_box[2]+1)
                    self.assertGreaterEqual(picture.top,pic_box[1]-1)
                    self.assertLessEqual(picture.top+picture.height+Inches(.059),caption.top)
                    width,height=picture.image.size
                    self.assertAlmostEqual(picture.width/picture.height,width/height,places=4)

    def test_long_intro_160_character_captions_and_three_240_character_requirements_fit_every_layout(self):
        prs=Presentation(TEMPLATE_FILE)
        summary=('完整问题表现说明须全部保留。'*20)[:160]
        intro=('本月重点检查概述内容须完整保留。'*20)[:160]
        issues=[{'issue_id':i,'station_name':'较长示例站名加油站','description':summary,
            'caption':summary,'issue_photo':'wide.jpg','requirement':(f'规范{i}须保持设施完好。'*30)[:240],
            'requirement_status':'ready'} for i in range(1,4)]
        for index in range(12,24):
            with self.subTest(slide=index+1):
                slide=prs.slides[index]
                evidence_page(slide,'2. 设备设施类问题',issues,{},issues,intro,self.root)
                heading=named(slide,'safety-evidence-intro')[0]
                norm=named(slide,'safety-evidence-requirements')[0]
                self.assertEqual(heading.text,intro)
                self.assertLessEqual(norm.top+norm.height,prs.slide_height)
                p=heading.text_frame.paragraphs[0]
                self.assertGreaterEqual(heading.height/914400+.001,
                    evidence_text_height([intro],heading.width/914400,p.runs[0].font.size.pt,p.line_spacing))
                for issue,caption,picture in zip(issues,named(slide,'safety-evidence-caption-'),named(slide,'safety-evidence-photo-')):
                    self.assertGreater(picture.height,0)
                    self.assertGreaterEqual(picture.top,heading.top+heading.height+Inches(.059))
                    self.assertLessEqual(picture.top+picture.height+Inches(.059),caption.top)
                    self.assertLessEqual(caption.top+caption.height+Inches(.079),norm.top)
                    paragraphs=caption.text_frame.paragraphs
                    self.assertIn(summary,[p.text for p in paragraphs])
                    self.assertIn(issue['requirement'],norm.text)
                    size=paragraphs[0].runs[0].font.size.pt
                    self.assertLessEqual(size,16)
                    self.assertGreaterEqual(caption.height/914400+.001,
                        evidence_text_height([p.text for p in paragraphs],caption.width/914400,size,1,6))
                size=norm.text_frame.paragraphs[0].runs[0].font.size.pt
                leading=norm.text_frame.paragraphs[0].line_spacing
                self.assertGreaterEqual(norm.height/914400+.001,
                    evidence_text_height([p.text for p in norm.text_frame.paragraphs],norm.width/914400,size,leading))

    def test_one_shared_red_frame_and_numbered_requirements_match_actual_photo_order(self):
        lookup={row['issue_id']:row for row in self.report['safety_issues']}
        for index in range(12,24):
            with self.subTest(slide=index+1):
                slide=self.prs.slides[index]
                norms=[s for s in slide.shapes if s.has_text_frame and s.text.startswith('规范要求')]
                self.assertEqual(len(norms),1)
                norm=norms[0]
                self.assertGreater(norm.width,Inches(12))
                self.assertEqual(str(norm.line.color.rgb),'C00000')
                self.assertEqual(norm.line.dash_style,MSO_LINE_DASH_STYLE.LONG_DASH)
                self.assertEqual(norm.line.width,Pt(1.5 if index>=20 else 1))
                self.assertTrue(norm._element.xpath('./p:spPr/a:noFill'))
                ids=photo_ids(slide)
                self.assertEqual(len(ids),3)
                for number,(paragraph,issue_id) in enumerate(zip(norm.text_frame.paragraphs,ids),1):
                    issue=lookup[issue_id]
                    expected=issue['requirement'] if issue.get('requirement_status') in ('ready','manual') else '暂无已确认的规范要求。'
                    self.assertEqual(paragraph.text,('规范要求：' if number==1 else '')+f'{number}. '+expected)
                    self.assertTrue(all(r.font.bold and str(r.font.color.rgb)=='C00000' for r in paragraph.runs))
                    self.assertEqual(paragraph.line_spacing,1.5)
                self.assertNotIn('旧评分',norm.text)

    def test_ranked_examples_monthly_distinct_classes_and_two_page_round_robin(self):
        self.assertEqual(photo_ids(self.prs.slides[12]),[3,5,6])
        self.assertEqual(photo_ids(self.prs.slides[13]),[3,5,6])
        self.assertEqual(photo_ids(self.prs.slides[14]),[1,4,2])
        self.assertEqual(photo_ids(self.prs.slides[18])+photo_ids(self.prs.slides[19]),[21,20,19,18,17,16])
        report=copy.deepcopy(self.report)
        report['safety_analysis']['onsite:topic:设备设施']['groups']=[
            group('首类',[1,2,3,4],[4,3,2,1]),group('次类',[5,6],[6,5])]
        with patch('safety_report_presentation.evidence_page',wraps=evidence_page) as render:
            build_safety_presentation(report,self.root/'sparse-monthly.pptx',self.root)
        monthly=next(call for call in render.call_args_list if call.args[1]=='1.月度重点工作类问题')
        self.assertEqual([row['issue_id'] for row in monthly.args[4]],[4,6])
        prs=Presentation(self.root/'sparse-monthly.pptx')
        self.assertEqual(photo_ids(prs.slides[12]),[4,6])
        self.assertEqual(len(named(prs.slides[12],'safety-evidence-caption-')),2)

    def test_monthly_intro_names_only_the_first_three_groups(self):
        report=copy.deepcopy(self.report)
        report['safety_analysis']['onsite:topic:设备设施']['groups']=[group('第一类',[1,2]),
            group('第二类',[3,4]),group('第三类',[5]),group('第四类不应进入概述',[6])]
        output=self.root/'three-monthly-classes.pptx'
        build_safety_presentation(report,output,self.root)
        intro=named(Presentation(output).slides[12],'safety-evidence-intro')[0].text
        for name in ('第一类','第二类','第三类'):self.assertIn(name,intro)
        self.assertNotIn('第四类',intro)

    def test_requirements_only_use_authoritative_ready_or_manual_snapshot_fields(self):
        prs=Presentation(TEMPLATE_FILE);slide=prs.slides[13]
        issues=[{'issue_id':i,'station_name':f'站{i}','description':f'现象{i}','issue_photo':'wide.jpg',
            'requirement':f'旧评分条款{i}扣10分。'} for i in range(1,4)]
        issues[0].update(requirement='实际规范应保持有效。',requirement_status='manual')
        issues[1]['requirement_status']='pending'
        chosen=[dict(row,requirement='不可使用的选图副本',requirement_status='ready') for row in reversed(issues)]
        evidence_page(slide,'2. 设备设施类问题',issues,{},chosen,'实际检查概述。',self.root)
        norm=named(slide,'safety-evidence-requirements')[0]
        self.assertEqual([p.text for p in norm.text_frame.paragraphs],[
            '规范要求：1. 暂无已确认的规范要求。','2. 暂无已确认的规范要求。','3. 实际规范应保持有效。'])
        self.assertNotIn('旧评分',norm.text)
        self.assertNotIn('副本',norm.text)
        self.assertNotIn('旧评分',slide.notes_slide.notes_text_frame.text)

    def test_long_confirmed_requirements_reflow_instead_of_truncating(self):
        prs=Presentation(TEMPLATE_FILE);slide=prs.slides[13]
        issues=[{'issue_id':i,'station_name':f'站{i}','description':'已核实的问题现象。','issue_photo':'wide.jpg',
            'requirement':f'规范{i}：'+('设施应保持完好并定期核查。'*18), 'requirement_status':'ready'} for i in range(1,4)]
        evidence_page(slide,'2. 设备设施类问题',issues,{},issues,'实际检查概述。',self.root)
        norm=named(slide,'safety-evidence-requirements')[0]
        self.assertEqual(len(norm.text_frame.paragraphs),3)
        self.assertGreater(norm.height,Inches(.9591))
        for issue,p in zip(issues,norm.text_frame.paragraphs):
            self.assertIn(issue['requirement'],p.text)
            self.assertNotIn('…',p.text)
        self.assertTrue(all(s.top+s.height<norm.top for s in named(slide,'safety-evidence-caption-')))
        self.assertTrue(all(s.height>0 for s in named(slide,'safety-evidence-photo-')))

    def test_backlog_preserves_native_chart_styles_geometry_and_red_integer_figures(self):
        for index,mode in [(27,'onsite'),(28,'video')]:
            with self.subTest(mode=mode):
                slide=self.prs.slides[index];source=self.template.slides[index]
                self.assertFalse(any(s.has_table for s in slide.shapes))
                graph=next(s for s in slide.shapes if s.has_chart)
                original=next(s for s in source.shapes if s.has_chart)
                self.assertEqual(geometry(graph),geometry(original))
                self.assertEqual(graph.chart.chart_type,XL_CHART_TYPE.COLUMN_CLUSTERED)
                chart=graph.chart;root=chart._chartSpace
                self.assertEqual(chart.chart_title.text_frame.text,original.chart.chart_title.text_frame.text)
                self.assertEqual([s.name for s in chart.series],['待签名','待整改'])
                self.assertEqual(root.xpath('.//c:title//a:defRPr')[0].get('sz'),'1680')
                self.assertEqual(root.xpath('.//c:legend/c:txPr//a:defRPr')[0].get('sz'),'1400')
                self.assertEqual(root.xpath('.//c:legend/c:legendPos')[0].get('val'),'b')
                self.assertEqual([n.get('sz') for n in root.xpath('.//c:ser/c:dLbls/c:txPr//a:defRPr')],['1400','1400'])
                self.assertEqual([n.get('val') for n in root.xpath('.//c:ser/c:spPr/a:solidFill/a:schemeClr')],['accent1','accent2'])
                self.assertEqual(root.xpath('.//c:barChart/c:gapWidth')[0].get('val'),'246')
                self.assertEqual(root.xpath('.//c:barChart/c:overlap')[0].get('val'),'-28')
                self.assertFalse(root.xpath('.//c:dTable'))
                self.assertTrue(all(n.get('formatCode')=='0' for n in root.xpath('.//c:numFmt')))
                self.assertTrue(all('.' not in (n.text or '') for n in root.xpath('.//c:numCache/c:pt/c:v')))
                self.assertIsNotNone(chart.part.chart_workbook.xlsx_part)
                body=next(s for s in slide.shapes if s.has_text_frame and s.text.startswith(('（一）','（二）')))
                old_body=next(s for s in source.shapes if s.has_text_frame and s.text.startswith(('（一）','（二）')))
                self.assertEqual((body.left,body.top,body.width),(old_body.left,old_body.top,old_body.width))
                self.assertLess(body.top+body.height,graph.top)
                self.assertEqual(len(body.text_frame.paragraphs),2)
                heading,narrative=body.text_frame.paragraphs
                for p in (heading,narrative):
                    props=p._p.find(qn('a:pPr'))
                    self.assertEqual((props.get('marL'),props.get('indent')),('0','542925'))
                    self.assertIsNotNone(props.find(qn('a:buNone')))
                self.assertEqual(heading.line_spacing,Pt(28))
                self.assertTrue(all(r.font.size==Pt(16) for r in heading.runs))
                self.assertIn('2026-09-30',heading.text)
                rows=sorted(self.report['pending_rectification'][mode],key=lambda r:unit_key(r['unit_name']))
                red=[r.text for r in narrative.runs if str(r.font.color.rgb)=='FF0000']
                self.assertEqual(red,[str(r[key]) for r in rows for key in ('pending_sign','pending_rectification')])
                self.assertEqual([list(s.values) for s in chart.series],
                    [[r[key] for r in rows] for key in ('pending_sign','pending_rectification')])

    def test_empty_backlog_and_empty_evidence_keep_native_objects_without_label_errors(self):
        prs=Presentation(TEMPLATE_FILE)
        for index,mode in [(27,'onsite'),(28,'video')]:
            slide=prs.slides[index]
            backlog_page(slide,[],mode,'2026-09-30')
            graph=next(s for s in slide.shapes if s.has_chart)
            self.assertEqual([list(s.values) for s in graph.chart.series],[[0],[0]])
            self.assertFalse(graph.chart.plots[0].data_labels.show_value)
            self.assertFalse(any(s.has_table for s in slide.shapes))
            self.assertIn('暂无待签名或待整改问题',''.join(s.text for s in slide.shapes if s.has_text_frame))
            backlog_page(slide,[{'unit_name':'浦东','pending_sign':12,'pending_rectification':0}],mode,'2026-09-30')
            self.assertIsNone(graph.chart.value_axis.maximum_scale)
            self.assertEqual(list(graph.chart.series[0].values),[12])
        slide=prs.slides[12]
        evidence_page(slide,'1.月度重点工作类问题',[],{},[],'当前暂无问题。',self.root)
        self.assertEqual(len(named(slide,'safety-evidence-requirements')),1)
        self.assertFalse(named(slide,'safety-evidence-photo-'))
        self.assertFalse(named(slide,'safety-evidence-caption-'))
        output=BytesIO();prs.save(output)
        self.assertEqual(len(Presentation(BytesIO(output.getvalue())).slides),31)

    def test_18_unit_backlog_fits_above_the_unchanged_native_chart(self):
        prs=Presentation(TEMPLATE_FILE)
        rows=[{'unit_name':name,'pending_sign':101+i,'pending_rectification':201+i}
            for i,name in enumerate(UNIT_ORDER[:18])]
        self.assertEqual(len(rows),18)
        for index,mode in [(27,'onsite'),(28,'video')]:
            slide=prs.slides[index];graph=next(s for s in slide.shapes if s.has_chart)
            box=geometry(graph)
            backlog_page(slide,rows,mode,'2026-09-30')
            self.assertEqual(geometry(graph),box)
            body=next(s for s in slide.shapes if s.has_text_frame and s.text.startswith(('（一）','（二）')))
            heading,narrative=body.text_frame.paragraphs
            for row in rows:
                self.assertIn(unit_label(row['unit_name']),narrative.text)
                self.assertIn(f"待签名{row['pending_sign']}项、待整改{row['pending_rectification']}项",narrative.text)
            self.assertNotIn('…',body.text)
            width=body.width/914400-.59375
            heights=[evidence_text_height([p.text],width,p.runs[0].font.size.pt,line_height=p.line_spacing.pt)
                for p in (heading,narrative)]
            self.assertLessEqual(sum(heights)-.1,body.height/914400)
            self.assertLess(body.top+body.height,graph.top)
            self.assertFalse(any(s.has_table for s in slide.shapes))

    def test_completely_empty_saved_report_still_builds_all_31_slides_without_ai(self):
        report=fixture_report(rows=[])
        output=self.root/'empty-all31.pptx'
        with patch('safety_report_analysis.generate_analysis',side_effect=AssertionError('No paid AI')), patch(
                'safety_report_examples.generate_representatives',side_effect=AssertionError('No paid AI')):
            build_safety_presentation(report,output,self.root)
        self.assertEqual(len(Presentation(output).slides),31)


if __name__=='__main__':
    unittest.main()
