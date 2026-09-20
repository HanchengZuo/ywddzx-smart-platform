import tempfile
import importlib.util
import os
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch
from datetime import date
from contextlib import ExitStack

from pptx import Presentation
from equipment_report_presentation import TEMPLATE_FILE, build_equipment_template_presentation, fill_station_page, normalize_template_fonts
from equipment_report_library import save_exclusions, exclusions
import app as reports


class EquipmentTemplateTest(unittest.TestCase):
    def test_native_template_preserves_fixed_pages_and_updates_data(self):
        report = {'month': '2026-09', 'summary': {'total_issue_count': 12, 'station_count': 3},
                  'region_rows': [{'unit_name': '浦东', 'unit_type': 'region', 'issue_count': 12,
                                   'station_count': 3, 'average_issue_count': 4}],
                  'station_ranking': [{'station_name': '甲站', 'issue_count': 8},
                                      {'station_name': '乙站', 'issue_count': 4},
                                      {'station_name': '丙站', 'issue_count': 0}],
                  'area_distribution': [{'name': '油罐区', 'count': 12, 'percentage': 100}],
                  'item_distribution': [{'name': '卸油口', 'count': 12, 'percentage': 100}]}
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'report.pptx'
            build_equipment_template_presentation(report, path)
            original, generated = Presentation(TEMPLATE_FILE), Presentation(path)
            normalize_template_fonts(original)
            self.assertEqual(len(generated.slides), 40)
            for index in [1, 2, 4, 7, *range(10, 40)]:
                self.assertEqual(original.slides[index]._element.xml, generated.slides[index]._element.xml)
            texts = lambda i: '\n'.join(s.text for s in generated.slides[i].shapes if s.has_text_frame)
            self.assertIn('2026年9月', texts(0))
            self.assertNotIn('199', texts(3))
            self.assertIn('4.0', texts(3))
            self.assertIn('1个片区、0家股权单位、共3座加油站', texts(5))
            for index, category in [(8, '油罐区'), (9, '卸油口')]:
                chart = next(s.chart for s in generated.slides[index].shapes if s.has_chart)
                self.assertEqual(chart.series[0].values, (12.0,))
                self.assertEqual(chart.plots[0].categories[0].label, category)
            table = next(s.table for s in generated.slides[6].shapes if s.has_table)
            self.assertEqual([row.cells[1].text for row in list(table.rows)[1:]], ['8', '4', '0'])

    def test_station_columns_fit_canvas_for_all_sizes(self):
        for count in [0, 1, 11, 12, 22, 23, 33, 34, 44, 45, 55]:
            with self.subTest(count=count):
                prs = Presentation(TEMPLATE_FILE)
                slide = prs.slides[6]
                fill_station_page(slide, [{'station_name': f'测试站{i}', 'issue_count': i} for i in range(count)])
                tables = [s for s in slide.shapes if s.has_table]
                self.assertEqual(sum(len(s.table.rows)-1 for s in tables), max(1, count))
                for shape in tables:
                    self.assertLessEqual(shape.top+shape.height, prs.slide_height)
                    self.assertLessEqual(shape.left+shape.width, prs.slide_width)

    def test_large_station_list_continues_without_losing_fixed_slides(self):
        report = {'month': '2026-09', 'summary': {}, 'station_ranking': [
            {'station_name': f'站{i}', 'issue_count': i} for i in range(60)
        ]}
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'report.pptx'
            build_equipment_template_presentation(report, path)
            prs = Presentation(path)
            self.assertEqual(len(prs.slides), 41)
            rows = [row.cells[1].text for index in [6, 7]
                    for shape in sorted([s for s in prs.slides[index].shapes if s.has_table], key=lambda s: s.left)
                    for row in list(shape.table.rows)[1:]]
            self.assertEqual(rows, list(map(str, range(59, -1, -1))))
            original = Presentation(TEMPLATE_FILE)
            normalize_template_fonts(original)
            self.assertEqual(prs.slides[8]._element.xml, original.slides[7]._element.xml)

    def test_exclusions_reject_out_of_scope_and_non_integer_ids(self):
        with patch('equipment_report_library.available', return_value=True):
            for values in [[99], ['1'], [True], None]:
                with self.assertRaises(ValueError):
                    save_exclusions(MagicMock(), [1, 2], values, 3)

    def test_library_query_requires_approved_completed_and_scope(self):
        cursor = MagicMock()
        cursor.fetchall.return_value = []
        with patch('app.append_inspection_table_scope_filter', return_value=True), patch(
            'app.append_station_region_scope_filter', return_value=True
        ):
            reports.fetch_equipment_report_issue_library(cursor, {'id': 3}, date(2026, 9, 1), date(2026, 10, 1))
        query = cursor.execute.call_args.args[0]
        for text in ['approved', '已确认完成', 'inspection_table_id', 'standard_detail_text']:
            self.assertIn(text, query)

    def test_issue_fields_and_zero_issue_stations_are_counted(self):
        report, _ = reports.build_equipment_facilities_report_payload(date(2026, 9, 1), [
            {'id': 5, 'station_id': 1, 'station_name': '甲站', 'description': '测试',
             'standard_detail_text': '所属区域：油罐区\n检查事项：卸油口'}
        ], [{'inspection_id': i, 'station_id': i, 'station_name': f'站{i}'} for i in [1, 2]])
        self.assertEqual(report['summary']['station_count'], 2)
        self.assertEqual(report['summary']['average_issue_count'], .5)
        self.assertEqual(report['area_distribution'][0]['name'], '油罐区')
        self.assertEqual(report['item_distribution'][0]['name'], '卸油口')

    def test_library_save_uses_server_actor_and_does_not_generate(self):
        with reports.app.test_request_context(method='PUT', json={
            'month': '2026-09', 'excluded_issue_ids': [5], 'user_id': 999
        }), patch('app.get_db_connection', return_value=MagicMock()), patch(
            'app.get_authorized_inspection_report_user', return_value={'id': 3}
        ), patch('app.fetch_equipment_report_issue_library', return_value=[{
            'id': 5, 'table_name': '设备设施检查表（现场）', 'standard_detail_text': '所属区域：油罐区'
        }]), patch('app.save_equipment_issue_exclusions', return_value={5}) as save, patch(
            'app.save_report_workspace'
        ) as workspace, patch('app.start_inspection_report_generation_job') as generate:
            response = reports.manage_equipment_report_issue_selection().get_json()
        self.assertTrue(response['success'])
        self.assertFalse(response['issues'][0]['included'])
        self.assertEqual(save.call_args.args[-1], 3)
        self.assertEqual(workspace.call_args.kwargs['section'], 'issue_library')
        generate.assert_not_called()

    def test_library_rejects_missing_permission(self):
        with reports.app.test_request_context(query_string={'month': '2026-09'}), patch(
            'app.get_db_connection', return_value=MagicMock()
        ), patch('app.get_authorized_inspection_report_user', side_effect=PermissionError('无权查看')):
            self.assertEqual(reports.manage_equipment_report_issue_selection()[1], 403)

    def test_generation_applies_saved_exclusions_without_ai(self):
        connection = MagicMock()
        connection.cursor.return_value.fetchall.side_effect = [
            [{'id': 5, 'station_id': 1, 'station_name': '甲站',
              'standard_detail_text': '所属区域：油罐区\n检查事项：卸油口'}],
            [{'inspection_id': 1, 'station_id': 1, 'station_name': '甲站'}]
        ]
        with ExitStack() as stack:
            for name, value in {
                'get_db_connection': connection, 'get_user_by_id': {'id': 3},
                'has_permission': True, 'append_inspection_table_scope_filter': True,
                'append_station_region_scope_filter': True,
                'resolve_inspection_report_source_selection': (None, {'mode': 'all'}),
                'equipment_issue_exclusions': {5}, 'update_inspection_report_job': None,
                'equipment_analysis.load_overrides': {},
            }.items():
                stack.enter_context(patch(f'app.{name}', return_value=value))
            ai = stack.enter_context(patch('app.generate_equipment_facilities_report_insights'))
            topics = stack.enter_context(patch('equipment_report_analysis.choose_topics'))
            save = stack.enter_context(patch('app.save_inspection_report_snapshot'))
            reports.generate_equipment_facilities_report_job('test', 3, '2026-09', 'shared', {})
        saved = save.call_args.args[4]
        self.assertEqual(saved['summary']['total_issue_count'], 0)
        self.assertEqual(saved['summary']['station_count'], 1)
        self.assertFalse(saved['issue_library_snapshot'][0]['included'])
        self.assertEqual(saved['area_distribution'], [])
        ai.assert_not_called()
        topics.assert_not_called()

    @unittest.skipUnless(os.getenv('ISSUE_LIFECYCLE_DB_TEST') == '1', 'isolated database test opt-in')
    def test_migration_and_selection_persistence(self):
        from psycopg2.extras import RealDictCursor
        import uuid
        connection = reports.get_db_connection()
        try:
            with connection.cursor(cursor_factory=RealDictCursor) as cursor:
                schema = 'equipment_test_' + uuid.uuid4().hex
                cursor.execute(f'CREATE SCHEMA {schema}')
                cursor.execute(f'SET LOCAL search_path TO {schema}')
                cursor.execute('CREATE TABLE issues(id bigint PRIMARY KEY)')
                cursor.execute('INSERT INTO issues VALUES(1),(2)')
                path = Path(__file__).parents[1] / 'migrations/versions/20260916_001_equipment_report_selection.py'
                spec = importlib.util.spec_from_file_location('equipment_migration', path)
                migration = importlib.util.module_from_spec(spec)
                spec.loader.exec_module(migration)
                with patch.object(migration.op, 'execute', side_effect=cursor.execute):
                    migration.upgrade()
                    save_exclusions(cursor, [1, 2], [1], 3)
                    migration.upgrade()
                    self.assertEqual(exclusions(cursor, [1, 2]), {1})
                    migration.downgrade()
                    self.assertEqual(exclusions(cursor, [1, 2]), {1})
                import equipment_report_analysis as topics
                path = Path(__file__).parents[1] / 'migrations/versions/20260920_001_equipment_topics.py'
                spec = importlib.util.spec_from_file_location('topics_migration', path)
                migration = importlib.util.module_from_spec(spec)
                spec.loader.exec_module(migration)
                with patch.object(migration.op, 'execute', side_effect=cursor.execute):
                    migration.upgrade()
                    topics.save_overrides(cursor, [{'issue_id':1}, {'issue_id':2}], {},
                        {'special_issue_ids':[1], 'severe_issue_ids':[2]}, 3)
                    migration.upgrade()
                    migration.downgrade()
                    saved = topics.load_overrides(cursor, [1,2])
                    self.assertTrue(saved[('special',1)])
                    self.assertFalse(saved[('severe',1)])
                cursor.execute('DELETE FROM issues WHERE id=1')
                self.assertNotIn(('special',1), topics.load_overrides(cursor, [1,2]))
                self.assertEqual(exclusions(cursor, [1, 2]), set())
                save_exclusions(cursor, [2], [2], 4)
                save_exclusions(cursor, [2], [], 5)
                self.assertEqual(exclusions(cursor, [2]), set())
        finally:
            connection.rollback()
            connection.close()


if __name__ == '__main__':
    unittest.main()
