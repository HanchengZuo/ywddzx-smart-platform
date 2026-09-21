import tempfile
import unittest
import zipfile
import warnings
from pathlib import Path
from unittest.mock import MagicMock, patch
from pptx import Presentation
import equipment_report_analysis as analysis
from equipment_report_presentation import build_equipment_template_presentation, TEMPLATE_FILE, normalize_template_fonts
from report_ai_memory import valid_insights


def issues():
    return [dict(issue_id=i, external_standard_id=9000 if i<4 else 9005,
                 station_id=i, station_name=f'站{i}', management_unit='浦东片区',
                 description=f'实际问题{i}', issue_photo='') for i in range(1,6)]


class EquipmentAnalysisTest(unittest.TestCase):
    def test_more_than_forty_slides_have_unique_package_parts(self):
        report = {'month':'2026-09', 'summary':{'station_count':60, 'total_issue_count':0},
                  'equipment_analysis':{'issues':[], 'phrase_distribution':[]},
                  'region_rows':[dict(unit_name='浦东',unit_type='region',station_count=60,issue_count=0,average_issue_count=0)],
                  'station_ranking':[dict(station_id=i,station_name=f'站{i}',management_unit='浦东',issue_count=0) for i in range(60)]}
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'large.pptx'
            with warnings.catch_warnings():
                warnings.simplefilter('error', UserWarning)
                build_equipment_template_presentation(report,path)
            with zipfile.ZipFile(path) as archive:
                names = archive.namelist()
                self.assertEqual(len(names), len(set(names)))
            prs = Presentation(path)
            self.assertGreater(len(prs.slides), 40)
            self.assertEqual(len({slide.part.partname for slide in prs.slides}), len(prs.slides))
            original = Presentation(TEMPLATE_FILE); normalize_template_fonts(original)
            self.assertEqual(prs.slides[-1]._element.xml,original.slides[-1]._element.xml)

    def test_ai_ids_normalized_without_inventing_references(self):
        context = {'kind': 'severe', 'issues': [{'issue_id': i} for i in range(1,6)]}
        self.assertEqual(analysis.normalize_topic_ids({'issue_ids':[' 1 ',1,'2',3,4]}, context), ([1,2,3], False))
        self.assertEqual(analysis.normalize_topic_ids({'issue_ids':[True,1.0,'1.0',9000,{'issue_id':1},'2']}, context), ([2], True))
        self.assertEqual(analysis.normalize_topic_ids({'issue_ids':[]}, context), ([], False))
        self.assertEqual(analysis.normalize_topic_ids({'issue_ids':[999]}, context), ([], True))
        self.assertEqual(analysis.normalize_topic_ids(None, context), ([], True))

    def test_invalid_ai_output_does_not_abort_or_become_successful_memory(self):
        client = MagicMock()
        response = client.with_options.return_value.chat.completions.create.return_value
        response.choices[0].message.content = '{"issue_ids":[99999]}'
        with patch('ai_utils.get_deepseek_client', return_value=client):
            result = analysis.choose_topics({'kind':'severe', 'issues':[{'issue_id':1}]})
        self.assertEqual(result['payload']['issue_ids'], [])
        self.assertFalse(result['generated'])
        self.assertTrue(result['usage']['fallback_used'])
        self.assertIn('warning', result)
        client.with_options.return_value.chat.completions.create.assert_called_once()

    def test_failed_selection_is_visible_and_manual_choices_still_apply(self):
        with patch.object(analysis, 'choose_topics', return_value={
            'payload':{'issue_ids':[]}, 'warning':'待核查'}):
            result = analysis.analyze(issues(), {('severe',1):True})
        self.assertEqual(result['severe_issue_ids'], [1])
        self.assertTrue(result['selection_warnings'])

    def test_reference_causes_are_not_generated(self):
        rows = analysis.enrich(issues())
        self.assertEqual(rows[0]['phrase'], '油罐区静电接地检测箱未接地')
        self.assertEqual(rows[0]['cause'], analysis.catalog()['9000']['cause'])
        self.assertEqual(rows[0]['management_unit'], '浦东管理片区')
        missing = analysis.enrich([dict(issue_id=99, external_standard_id=1)])[0]
        self.assertIn('未匹配', missing['phrase'])
        self.assertEqual(missing['cause'], '参考表未提供原因')

    def test_high_group_evidence_complete_and_special_excludes_high(self):
        calls = []
        def choose(context):
            calls.append(context)
            return {'payload': {'issue_ids': [context['issues'][0]['issue_id']]}}
        with patch.object(analysis, 'choose_topics', side_effect=choose):
            result = analysis.analyze(issues())
        self.assertEqual([i['issue_id'] for i in result['high_groups'][0]['issues']], [1,2,3])
        self.assertEqual(result['special_issue_ids'], [4])
        self.assertTrue(all(i['issue_id'] not in [1,2,3] for c in calls if c['kind']=='special' for i in c['issues']))
        self.assertEqual(sum(row['count'] for row in result['phrase_distribution']), 5)
        self.assertEqual(result['phrase_distribution'][0]['percentage'], 60)

    def test_manual_overrides_and_scope_validation(self):
        rows = analysis.enrich(issues())
        base = {'high_groups': [analysis.groups_for(rows)[0]], 'special_issue_ids':[4], 'severe_issue_ids':[1]}
        result = analysis.apply_overrides(base, rows, {('special',4):False,('special',5):True,('severe',2):True})
        self.assertEqual(result['special_issue_ids'], [5])
        self.assertEqual(result['severe_issue_ids'], [1,2])
        cur = MagicMock(); cur.fetchone.return_value = {'name':'exists'}
        for data in [{'special_issue_ids':[1], 'severe_issue_ids':[]},
                     {'special_issue_ids':[], 'severe_issue_ids':[999]},
                     {'special_issue_ids':[], 'severe_issue_ids':[True]}]:
            with self.assertRaises(ValueError):
                analysis.save_overrides(cur, rows, base, data, 3)

    def test_stable_inputs_for_repeated_generation(self):
        calls = []
        def choose(context):
            calls.append(context)
            return {'payload': {'issue_ids':[]}}
        with patch.object(analysis, 'choose_topics', side_effect=choose):
            analysis.analyze(issues())
            first = list(calls); calls.clear()
            analysis.analyze(list(reversed(issues())))
        self.assertEqual(first, calls)
        self.assertTrue(valid_insights({'issue_ids':[]}, 'equipment_topics', {1}))
        self.assertFalse(valid_insights({'issue_ids':[99]}, 'equipment_topics', {1}))

    def test_native_pages_cover_all_stations_and_final_template(self):
        with patch.object(analysis, 'choose_topics', side_effect=lambda c: {'payload': {'issue_ids':[c['issues'][0]['issue_id']]}}):
            result = analysis.analyze(issues())
        long_title = '加油机内各元件静电接地线是否规范连接至接地端子排不合规'
        result['high_groups'][0]['phrase'] = long_title
        result['phrase_distribution'] = [dict(name=r['phrase'],count=1,percentage=1.5) for r in analysis.catalog().values()]
        report = {'month':'2026-09', 'summary':{'station_count':6,'total_issue_count':5},
                  'equipment_analysis':result,
                  'region_rows':[dict(unit_name='浦东',unit_type='region',station_count=6,issue_count=5,average_issue_count=.8)],
                  'station_ranking':[dict(station_id=i,station_name=f'站{i}',management_unit='浦东',issue_count=int(i<6)) for i in range(1,7)]}
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'report.pptx'
            build_equipment_template_presentation(report,path)
            prs = Presentation(path)
            original = Presentation(TEMPLATE_FILE); normalize_template_fonts(original)
            self.assertEqual(prs.slides[-1]._element.xml, original.slides[-1]._element.xml)
            content = '\n'.join(s.text for page in list(prs.slides)[10:] for s in page.shapes if s.has_text_frame)
            for i in range(1,7): self.assertIn(f'站{i}',content)
            self.assertIn('实际问题5',content)
            self.assertNotIn('龚路',content)
            self.assertNotIn('环南',content)
            self.assertEqual(len(prs.slides),21)
            page_texts = ['\n'.join(s.text for s in page.shapes if s.has_text_frame) for page in prs.slides]
            self.assertEqual(sum('加油站设备设施各类问题占比情况' in value for value in page_texts), 1)
            self.assertEqual(sum('加油站设备设施高频问题原因分析' in value for value in page_texts), 1)
            self.assertIn(long_title, page_texts[11].replace('\n',''))
            self.assertIn('检查站数：6站（站1、站2、站3、站4、站5、站6）', ''.join(page_texts).replace('\n',''))
            for page in list(prs.slides)[10:-1]:
                for shape in page.shapes:
                    self.assertLessEqual(shape.top+shape.height,prs.slide_height)
                    self.assertLessEqual(shape.left+shape.width,prs.slide_width)
