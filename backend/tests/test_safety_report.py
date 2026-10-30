import copy
import importlib.util
import os
from datetime import date
from pathlib import Path
import tempfile
import unittest
import uuid
from zipfile import ZipFile
from unittest.mock import MagicMock, patch

from flask import Flask
from PIL import Image
from pptx import Presentation
from psycopg2.extras import RealDictCursor
import psycopg2
import sqlalchemy as sa
from alembic.migration import MigrationContext
from alembic.operations import Operations

from safety_report_data import serialize, scope, period_label, pending_at, backlog, category_rows, focus_topics
from safety_report_analysis import (validate_analysis, analysis_context, normalize_settings,
    effective_analysis, build_report, photo_examples, load_settings, save_settings)
from safety_report_http import register_safety_report, config_categories
from safety_report_presentation import build_safety_presentation, TEMPLATE_FILE, SLIDE_ORDER
from safety_report_presentation import video_typical_narrative
from safety_template_styles import editable_plot


def raw_issue(issue_id=1, mode='onsite', **extra):
    return dict(dict(id=issue_id, station_id=1, station_name='测试站', region='浦东管理片区',
        checklist_mode='online' if mode == 'video' else 'offline', standard_id='1234',
        standard_detail_text='序号：26\n检查主题：设备设施\n检查内容：油罐操作井\n评分说明：跨接及接地规范，得10分。',
        description='油罐操作井接地连接缺失', issue_photo='', audit_status='approved',
        status='待整改'), **extra)


def mock_analysis(issues, *args):
    return {'groups': [{'name': '装置维护不到位', 'summary': '接地装置连接缺失，应按规范维护。',
        'issue_ids': [row['issue_id'] for row in issues]}] if issues else [],
        'causes': ['可能存在日常检查遗漏，需核实。'] if issues else [], 'suggestions': []}


def fixture_report(photo='', rows=None, settings=None):
    if rows is None:
        rows = [raw_issue(), raw_issue(2, 'video'), raw_issue(3)]
        rows[1]['standard_detail_text'] = '序号：2\n检查内容：加油站设备维修作业\n评分说明：执行检维修许可。'
    else: rows = copy.deepcopy(rows)
    for row in rows: row['issue_photo'] = photo
    def examples(context):
        return {'payload':{'selections':[{'name':g['name'],'issue_ids':[r['issue_id'] for r in g['issues']][:6]}
            for g in context['groups']]}}
    with patch('safety_report_examples.generate_representatives', side_effect=examples), patch('safety_report_analysis.analyze', side_effect=mock_analysis), patch(
        'safety_report_analysis.generate_analysis', return_value={'payload': {'groups': [], 'causes': [],
        'suggestions': [{'title': '强化设施维护', 'content': '开展接地连接专项检查，核实问题整改。'}]}}):
        report = build_report(date(2026, 9, 1), date(2026, 10, 1), rows, [],
            {'onsite': [{'unit_name': '浦东', 'pending_sign': 2, 'pending_rectification': 1}], 'video': []}, normalize_settings(settings or {}))
    report['summary'].update(date_from='2026-09-01', date_to='2026-09-30')
    return report


def focus_fixture_rows():
    rows = [raw_issue(i, standard_detail_text=f'检查主题：设备设施\n检查内容：设备子项{i}\n评分说明：设备规范') for i in range(1, 5)]
    rows += [raw_issue(i, standard_detail_text='检查主题：污染防治\n检查内容：环保设施配置齐全，有效运行。\n评分说明：环保规范') for i in range(5, 8)]
    rows += [raw_issue(i, 'video', standard_detail_text='检查内容：设备维修作业') for i in range(8, 16)]
    return rows


class SafetyDataTest(unittest.TestCase):
    def test_focus_counts_topics_not_content_or_video_and_preserves_exact_risk_topic(self):
        rows = [serialize(row) for row in focus_fixture_rows()]
        self.assertEqual(focus_topics(rows), [
            {'name': '设备设施', 'count': 4, 'percentage': 57.14},
            {'name': '污染防治', 'count': 3, 'percentage': 42.86}])
        self.assertEqual([row['issue_id'] for row in category_rows(rows, 'onsite', 'topic:设备设施')], [1, 2, 3, 4])
        risks = [dict(rows[0], topic='风险隐患'), dict(rows[1], topic='合规性管理')]
        self.assertEqual(len(category_rows(risks, 'onsite', 'topic:风险隐患')), 1)
        self.assertEqual(len(category_rows(risks, 'onsite', '风险合规')), 2)

    def test_content_aliases_and_ordinal_other(self):
        from safety_report_data import ALIASES
        for prefix, alias in ALIASES:
            row = raw_issue(); row['standard_detail_text'] = '检查主题：设备设施\n检查内容：'+prefix+'。\n评分说明：第一行\n第二行\n10分：完整达标'
            value = serialize(row)
            self.assertEqual(value['category_name'], alias)
            self.assertEqual(value['requirement'], '第一行\n第二行\n10分：完整达标')
        row['standard_detail_text'] = '序号：50\n类别：基础项\n检查主题：其他对生产经营有影响的问题隐患\n总分值：30\n评分说明：评分'
        self.assertEqual((serialize(row)['topic'], serialize(row)['category_name']), ('其他隐患类', '其他'))
        self.assertEqual(serialize(row)['requirement'], '评分')
        row = raw_issue(2, 'video'); row['standard_detail_text'] = '序号：13\n检查内容：原其他项目'
        self.assertEqual(serialize(row)['category_name'], '其他隐患类')

    def test_scope_separates_modes_and_counts_zero_issue_stations(self):
        visits = [dict(station_id=1, region='浦东管理片区', checklist_mode='offline'),
            dict(station_id=2, region='中油奉贤公司', checklist_mode='offline'),
            dict(station_id=3, region='宝静管理片区', checklist_mode='online')]
        value = scope([serialize(raw_issue()), serialize(raw_issue(2, 'video'))], visits, 'onsite', '9月')
        self.assertEqual((value['station_count'], value['total_issue_count'], value['average_issue_count']), (2, 1, .5))
        self.assertEqual((value['region_count'], value['holding_unit_count']), (1, 1))
        self.assertEqual(scope([], [], 'video', '')['average_issue_count'], 0)
        self.assertEqual(serialize(raw_issue(station_name='中油奉贤第一'))['unit_name'], '浦东')

    def test_dates_month_or_exact_cross_month(self):
        self.assertEqual(period_label(date(2026, 9, 4), date(2026, 10, 1)), '2026年09月')
        self.assertEqual(period_label(date(2026, 9, 1), date(2026, 11, 1)), '2026年09月01日至2026年10月31日')
        from safety_report_data import is_before
        self.assertFalse(is_before('2026-09-30T20:00:00+00:00', date(2026, 10, 1)))
        self.assertTrue(is_before('2026-09-30T15:59:59+00:00', date(2026, 10, 1)))

    def test_historical_pending_ignores_later_review_and_respects_reset(self):
        end = date(2026, 10, 1)
        row = raw_issue(); row.update(status='已闭环', audited_at='2026-09-01', station_manager_signed_at='2026-09-03')
        history = [dict(action_type='audit_changed', to_status='审核通过', created_at='2026-09-01'),
            dict(action_type='inspection_signed', to_status='站经理签名确认', created_at='2026-09-03'),
            dict(action_type='review_submitted', to_status='已闭环', created_at='2026-10-03')]
        self.assertEqual(pending_at(row, history, end), '待整改')
        history.append(dict(action_type='rectification_submitted', to_status='待复核', created_at='2026-09-20'))
        history.sort(key=lambda r: r['created_at'])
        self.assertIsNone(pending_at(row, history, end))
        history.append(dict(action_type='inspection_signed', to_status='撤销站经理签名', created_at='2026-09-25'))
        self.assertEqual(pending_at(row, [history[0], history[-1]], end), '待签名')
        self.assertEqual(backlog([row], [], end, [])['onsite'], [])
        self.assertIsNone(pending_at(dict(row, audit_status='rejected'), [], end))


class SafetyAnalysisTest(unittest.TestCase):
    def test_monthly_focus_uses_top_topic_while_slide_11_distribution_stays_content_based(self):
        report = fixture_report(rows=focus_fixture_rows())
        self.assertEqual(report['focus_category'], '设备设施')
        self.assertEqual(report['sections'][0]['detail_distribution'][0]['name'], '环保设施')
        selected = report['safety_analysis']['onsite:topic:设备设施']['groups'][0]['issue_ids']
        self.assertEqual(selected, [1, 2, 3, 4])
        self.assertEqual(report['safety_template_version'], 2)

    def test_manual_topic_focus_and_assignments_are_preserved(self):
        settings = {'focus_category': '污染防治', 'classifications': {'onsite:topic:污染防治': {
            'assignments': {'5': '人工环保分类'}, 'summaries': {'人工环保分类': '人工环保摘要'}}}}
        report = fixture_report(rows=focus_fixture_rows(), settings=settings)
        self.assertEqual(report['focus_category'], '污染防治')
        group = next(group for group in report['safety_analysis']['onsite:topic:污染防治']['groups'] if group['name'] == '人工环保分类')
        self.assertEqual((group['issue_ids'], group['summary']), ([5], '人工环保摘要'))
        self.assertEqual(report['safety_settings']['classifications'], settings['classifications'])

    def test_old_content_focus_falls_back_to_top_topic_without_deleting_saved_overrides(self):
        settings = {'focus_category': '环保设施', 'classifications': {'onsite:content:环保设施': {
            'assignments': {'5': '原人工分类'}, 'summaries': {}}}}
        report = fixture_report(rows=focus_fixture_rows(), settings=settings)
        self.assertEqual(report['focus_category'], '设备设施')
        self.assertEqual(report['safety_settings']['classifications'], settings['classifications'])
        self.assertNotIn('onsite:content:环保设施', report['safety_analysis'])

    def test_topic_validation_and_same_topic_memory_signatures(self):
        rows = [serialize(row) for row in focus_fixture_rows()]
        self.assertEqual(normalize_settings({'focus_category': '设备设施'}, rows)['focus_category'], '设备设施')
        with self.assertRaises(ValueError): normalize_settings({'focus_category': '环保设施'}, rows)
        selected = category_rows(rows, 'onsite', 'topic:设备设施')
        self.assertEqual(analysis_context(selected, 'onsite', 'topic:设备设施'), analysis_context(selected, 'onsite', '设备设施'))

    def test_ai_ids_complete_unique_and_no_unrelated_refs(self):
        issues = [serialize(raw_issue())]
        valid = mock_analysis(issues)
        self.assertEqual(validate_analysis(valid, issues), valid)
        for ids in ([], [1, 1], [999]):
            invalid = copy.deepcopy(valid); invalid['groups'][0]['issue_ids'] = ids
            with self.assertRaises(ValueError): validate_analysis(invalid, issues)
        invalid = copy.deepcopy(valid); invalid['causes'] = []
        with self.assertRaises(ValueError): validate_analysis(invalid, issues)

    def test_manual_grouping_summary_overrides_ai(self):
        rows = [serialize(raw_issue())]
        value = effective_analysis(mock_analysis(rows), rows,
            {'assignments': {'1': '人工核定分类'}, 'summaries': {'人工核定分类': '人工摘要'}})
        self.assertEqual(value['groups'], [{'name': '人工核定分类', 'issue_ids': [1], 'summary': '人工摘要'}])
        rows[0]['issue_photo'] = 'issue.jpg'
        self.assertEqual(photo_examples(rows, value, 3)[0]['caption'], '人工摘要')

    def test_memory_context_stable_without_dates_or_query_order(self):
        rows = [serialize(raw_issue(2)), serialize(raw_issue())]
        self.assertEqual(analysis_context(rows, 'onsite', '设备设施'), analysis_context(list(reversed(rows)), 'onsite', '设备设施'))
        self.assertNotIn('report_date', str(analysis_context(rows, 'onsite', '设备设施')))

    def test_invalid_saved_sources_and_overrides_rejected(self):
        for raw in ({'station_ids': ['1']}, {'typical_issue_ids': {'onsite': 2}},
            {'classifications': {'onsite:设备设施': {'assignments': {'999': '错'}}}}):
            with self.assertRaises(ValueError): normalize_settings(raw, [serialize(raw_issue())], {1})

    def test_empty_report_never_calls_ai(self):
        with patch('safety_report_analysis.generate_analysis', side_effect=AssertionError('paid AI')):
            report = build_report(date(2026, 9, 1), date(2026, 10, 1), [], [], {'onsite': [], 'video': []}, normalize_settings({}))
        self.assertEqual(report['summary']['total_issue_count'], 0)


class SafetyHttpTest(unittest.TestCase):
    def setUp(self):
        self.conn = MagicMock(); self.authorize = MagicMock(return_value={'id': 7, 'real_name': '测试保存人'})
        self.permission = MagicMock(return_value=False); self.workspace = MagicMock()
        self.app = Flask(__name__)
        register_safety_report(self.app, connect=lambda: self.conn, authorize=self.authorize,
            period=lambda *a: (date(2026, 9, 1), date(2026, 10, 1)),
            fetch_inputs=lambda *a: ([raw_issue()], [], []), has_permission=self.permission,
            save_workspace=self.workspace, log_usage=MagicMock())
        self.client = self.app.test_client()

    def test_configuration_focus_options_and_ai_scopes_use_topics(self):
        issues = [serialize(row) for row in focus_fixture_rows()]
        options, topics, focus = config_categories(issues, normalize_settings({}))
        self.assertEqual(focus, '设备设施')
        self.assertEqual([row['name'] for row in topics], ['设备设施', '污染防治'])
        self.assertIn(('onsite', 'topic:设备设施', '1. 月度重点：设备设施'), options)
        self.assertFalse(any(category.startswith('content:') for _, category, _ in options))
        with patch('safety_report_http.load_settings', return_value={'settings': {}, 'revision': 0}), patch(
            'safety_report_http.cached_analysis', return_value={}):
            config = self.client.get('/api/inspection-reports/safety-config').get_json()
            for value, status in [('设备设施', 200), ('油罐操作井', 400)]:
                with self.subTest(value=value), patch('safety_report_http.save_settings') as save:
                    response = self.client.put('/api/inspection-reports/safety-config', json={
                        'settings': {'focus_category': value}, 'revision': 0, 'data_signature': config['data_signature']})
                    self.assertEqual(response.status_code, status)
                    if status == 400: save.assert_not_called()

    def test_explicit_topic_reclassification_uses_only_eligible_topic_and_rejects_content_scope(self):
        from report_ai_memory import fingerprint
        self.permission.return_value = True
        with patch('safety_report_http.load_settings', return_value={'settings': {}, 'revision': 0}), patch(
            'safety_report_http.begin_report_ai_memory', return_value=None), patch(
            'safety_report_http.analyze', return_value={}) as analyze:
            request = {'mode': 'onsite', 'category': 'topic:设备设施', 'revision': 0,
                'data_signature': fingerprint([serialize(raw_issue())])}
            self.assertEqual(self.client.post('/api/inspection-reports/safety-analysis', json=request).status_code, 200)
            self.assertEqual(analyze.call_args.args[0][0]['topic'], '设备设施')
            analyze.reset_mock()
            request['category'] = 'content:油罐操作井'
            self.assertEqual(self.client.post('/api/inspection-reports/safety-analysis', json=request).status_code, 400)
            analyze.assert_not_called()

    def test_read_and_save_never_call_ai_and_preserve_server_identity(self):
        with patch('safety_report_http.load_settings', return_value={'settings': {}, 'revision': 0}), patch(
            'safety_report_http.cached_analysis', return_value={}), patch('safety_report_http.analyze', side_effect=AssertionError('AI')):
            result = self.client.get('/api/inspection-reports/safety-config').get_json()
            self.assertEqual(result['summary']['issue_count'], 1)
            with patch('safety_report_http.save_settings') as save:
                response = self.client.put('/api/inspection-reports/safety-config', json={
                    'settings': {}, 'revision': 0, 'data_signature': result['data_signature'], 'user_id': 999})
                self.assertEqual(response.status_code, 200)
                self.assertEqual(save.call_args.args[-1]['id'], 7)
                self.workspace.assert_called_once()

    def test_topic_analysis_read_is_readonly_permission_checked_and_never_calls_ai(self):
        from report_ai_memory import fingerprint
        analysis = {'groups': [], 'complete': False, 'classified_count': 0, 'issue_count': 1}
        with patch('safety_report_http.load_settings', return_value={'settings': {}, 'revision': 3}), patch(
            'safety_report_http.cached_analysis', return_value=analysis), patch('safety_report_http.analyze') as ai:
            result = self.client.get('/api/inspection-reports/safety-analysis?mode=onsite&category=topic:设备设施')
            self.assertEqual(result.status_code, 200)
            self.assertEqual(result.get_json()['revision'], 3)
            self.assertEqual(result.get_json()['data_signature'], fingerprint([serialize(raw_issue())]))
            self.assertFalse(result.get_json()['analysis']['complete'])
            self.assertEqual(len(result.get_json()['issues']), 1)
            self.assertEqual(self.client.get('/api/inspection-reports/safety-analysis?mode=onsite&category=topic:不存在').status_code, 400)
            ai.assert_not_called(); self.workspace.assert_not_called()
            self.permission.assert_not_called()  # Page authorization is enough for a free read.
        self.authorize.side_effect = PermissionError('无查看权限')
        self.assertEqual(self.client.get('/api/inspection-reports/safety-analysis').status_code, 403)

    def test_no_page_permission_and_no_generation_permission(self):
        self.authorize.side_effect = PermissionError('无查看权限')
        self.assertEqual(self.client.get('/api/inspection-reports/safety-config').status_code, 403)
        self.conn.rollback.assert_called()
        self.authorize.side_effect = None
        with patch('safety_report_http.load_settings', return_value={'settings': {}, 'revision': 0}), patch('safety_report_http.analyze') as ai:
            self.assertEqual(self.client.post('/api/inspection-reports/safety-analysis', json={}).status_code, 403)
            ai.assert_not_called()

    def test_failed_save_rolls_back_and_never_touches_workspace(self):
        from report_ai_memory import fingerprint
        with patch('safety_report_http.load_settings', return_value={'settings': {}, 'revision': 0}), patch(
            'safety_report_http.save_settings', side_effect=ValueError('配置已更新')):
            response = self.client.put('/api/inspection-reports/safety-config', json={
                'settings': {}, 'revision': 0, 'data_signature': fingerprint([serialize(raw_issue())])})
        self.assertEqual(response.status_code, 400)
        self.conn.rollback.assert_called(); self.conn.commit.assert_not_called(); self.workspace.assert_not_called()


class SafetyPresentationTest(unittest.TestCase):
    def test_overview_preserves_indented_paragraph_red_metrics_integer_chart_and_no_extra_table(self):
        from pptx.oxml.ns import qn
        with tempfile.TemporaryDirectory() as root:
            path = Path(root)/'report.pptx'
            report = fixture_report(); build_safety_presentation(report, path, root)
            prs = Presentation(path)
            for index in (2, 3):
                slide = prs.slides[index]
                body = next(s for s in slide.shapes if s.has_text_frame and s.text.startswith(('（一）', '（二）')))
                self.assertEqual(len(body.text_frame.paragraphs), 2)
                self.assertIn('各管理片区及合资公司问题数量如下：', body.text)
                self.assertIn('浦东片区', body.text)
                a, b = body.text_frame.paragraphs
                self.assertEqual(a._p.find(qn('a:pPr')).get('indent'), b._p.find(qn('a:pPr')).get('indent'))
                self.assertGreater(int(a._p.find(qn('a:pPr')).get('indent')), 0)
                red = [r.text for r in b.runs if str(r.font.color.rgb) == 'FF0000']
                self.assertTrue(any(text.startswith('，站均问题') and text.endswith('个') for text in red))
                self.assertEqual(len(red), 3)
                self.assertEqual(len({r.font.size for r in b.runs}), 1)
                self.assertFalse(any(s.has_table for s in slide.shapes))
                shape = next(s for s in slide.shapes if s.has_chart); chart = shape.chart
                self.assertIn('问题数量分布图', chart.chart_title.text_frame.text)
                self.assertEqual([s.name for s in chart.series], ['站点数量', '问题数量'])
                self.assertEqual([str(s.format.fill.fore_color.rgb) for s in chart.series], ['4F81BD', 'C0504D'])
                self.assertTrue(all(node.get('formatCode') == '0' for node in chart._chartSpace.xpath('.//c:dLbls/c:numFmt')))
                self.assertLess(body.top+body.height, shape.top)

    def test_distribution_reuses_native_3d_chart_with_titles_percentages_and_template_frames(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root)/'report.pptx';build_safety_presentation(fixture_report(), path, root)
            prs = Presentation(path)
            for index, mode in [(7, '现场'), (8, '视频')]:
                slide = prs.slides[index]
                text = '\n'.join(s.text for s in slide.shapes if s.has_text_frame)
                self.assertIn('本月检查共发现'+mode+'问题', text)
                self.assertIn('类问题：', text)
                chart = next(s.chart for s in slide.shapes if s.has_chart)
                self.assertEqual(len(chart._chartSpace.xpath('.//c:pie3DChart')), 1)
                self.assertEqual(len(chart._chartSpace.xpath('.//c:view3D')), 1)
                self.assertIn(mode+'巡查问题汇总分析',[s.text for s in slide.shapes if s.has_text_frame])
                self.assertFalse(chart._chartSpace.xpath('.//c:ser/c:dLbls'))
                labels = chart._chartSpace.xpath('.//c:pie3DChart/c:dLbls')[0]
                self.assertEqual(labels.find('{*}showVal').get('val'), '0')
                self.assertEqual(labels.find('{*}showPercent').get('val'), '1')
                with editable_plot(chart): self.assertTrue(chart.series[0].values)
                self.assertEqual(len(chart._chartSpace.xpath('.//c:pie3DChart')), 1)
                self.assertTrue(any(s.name.startswith('矩形') for s in slide.shapes))

    def test_typical_appears_once_with_frames_arrows_station_above_summary_and_no_video_station(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root)/'report.pptx'
            build_safety_presentation(fixture_report(settings={'typical_issue_ids': {'onsite': 1, 'video': 2}}), path, root)
            prs = Presentation(path)
            headings = [s.text for slide in prs.slides for s in slide.shapes if s.has_text_frame and s.top<914400]
            self.assertEqual(headings.count('二、检查发现-典型问题'), 1)
            slide = prs.slides[5]
            onsite = next(s for s in slide.shapes if s.name=='文本框 1')
            video = next(s for s in slide.shapes if s.name=='文本框 4')
            self.assertTrue(onsite.text.startswith('测试站加油站\n'))
            self.assertNotIn('测试站', video.text)
            self.assertTrue(video.text.startswith('9月视频扫站巡查过程中'))
            self.assertIn('总计出现1次', video.text)
            self.assertEqual(len([s for s in slide.shapes if s.name.startswith('矩形')]), 2)
            arrows=[s for s in slide.shapes if s.name.startswith('safety-split-arrow-')]
            self.assertEqual(len(arrows),4)
            self.assertEqual(sorted(str(s.fill.fore_color.rgb) for s in arrows),['000000','000000','FFFFFF','FFFFFF'])
            self.assertTrue(all(r.font.bold for r in onsite.text_frame.paragraphs[-1].runs))

    def test_video_typical_uses_only_the_highest_category_and_supports_saved_selection(self):
        report = {'month_label': '2026年09月', 'safety_issues': [{'issue_id': i,'mode': 'video'} for i in range(1,11)],
            'safety_analysis': {'video:__typical__': {'groups': [
                {'name': '最少表现', 'issue_ids': [10]}, {'name': '最多表现', 'issue_ids': [1,2,3,4]},
                {'name': '第二表现', 'issue_ids': [5,6,7]}, {'name': '第三表现', 'issue_ids': [8,9]},
                {'name': '无关问题', 'issue_ids': [999]}]}}}
        self.assertEqual(video_typical_narrative(report), '9月视频扫站巡查过程中，发现最多表现，总计出现4次。')
        report['safety_settings']={'typical_video_category':'第二表现'}
        self.assertEqual(video_typical_narrative(report), '9月视频扫站巡查过程中，发现第二表现，总计出现3次。')
        report['month_label'] = '2026年08月01日至2026年09月30日'
        self.assertTrue(video_typical_narrative(report).startswith(report['month_label']))

    def test_analysis_pages_have_two_red_square_sections_and_independent_native_3d_workbooks(self):
        with tempfile.TemporaryDirectory() as root:
            path=Path(root)/'report.pptx';build_safety_presentation(fixture_report(),path,root)
            prs=Presentation(path)
            for page,source in [(10,7),(11,8)]:
                slide=prs.slides[SLIDE_ORDER.index(page)]
                body=next(s for s in slide.shapes if s.has_text_frame and s.text.startswith('问题分布'))
                self.assertIn('问题分布：如右图所示',body.text)
                self.assertIn('产生原因：',body.text)
                self.assertNotIn('占比较高问题分析',body.text)
                self.assertEqual(len(body._element.xpath('.//a:buChar[@char="■"]')),2)
                self.assertTrue(all(n.get('val')=='FF0000' for n in body._element.xpath('.//a:buClr/a:srgbClr')))
                graphs=[s for s in slide.shapes if s.has_chart]
                self.assertEqual(len(graphs),1)
                self.assertTrue(graphs[0].chart._chartSpace.xpath('.//c:pie3DChart'))
                original=next(s for s in prs.slides[source].shapes if s.has_chart)
                self.assertIsNot(graphs[0].chart.part,original.chart.part)
                self.assertIsNot(graphs[0].chart.part.chart_workbook.xlsx_part,original.chart.part.chart_workbook.xlsx_part)

    def test_native_editable_static_pages_and_geometry(self):
        with tempfile.TemporaryDirectory() as root:
            Image.new('RGB', (150, 600), 'steelblue').save(Path(root)/'photo.jpg')
            report = fixture_report('photo.jpg')
            output = Path(root)/'safety.pptx'
            result = build_safety_presentation(report, output, root)
            prs = Presentation(output); template = Presentation(TEMPLATE_FILE)
            self.assertEqual(result['slide_count'], 31)
            for before, after in [(1, 1), (4, 4), (6, 6), (9, 9), (24, 24), (26, 26), (30, 30)]:
                self.assertEqual([s.text for s in template.slides[before].shapes if s.has_text_frame],
                    [s.text for s in prs.slides[SLIDE_ORDER.index(after)].shapes if s.has_text_frame])
            charts = [s.chart for slide in prs.slides for s in slide.shapes if s.has_chart]
            self.assertGreater(len(charts), 5)
            self.assertEqual(len([s for s in prs.slides[SLIDE_ORDER.index(11)].shapes if s.has_chart]), 1)
            self.assertTrue(any(s.shape_type == 13 for s in prs.slides[5].shapes))
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_chart or shape.has_table or (shape.has_text_frame and shape.top > 914400):
                        self.assertLessEqual(shape.left + shape.width, prs.slide_width + 1000)
                        self.assertLessEqual(shape.top + shape.height, prs.slide_height + 1000)
            with ZipFile(output) as z:
                self.assertFalse(any('comments/' in name for name in z.namelist()))
                self.assertTrue(any('embeddings/' in name for name in z.namelist()))
            self.assertNotIn('2021年4月', '\n'.join(s.text for slide in prs.slides for s in slide.shapes if s.has_text_frame))

    def test_old_snapshot_requires_regeneration_instead_of_stale_ppt(self):
        with self.assertRaises(ValueError): build_safety_presentation({}, '/tmp/should-not-exist-safety.pptx')
        old = fixture_report(); old['safety_template_version'] = 1
        with self.assertRaises(ValueError): build_safety_presentation(old, '/tmp/should-not-exist-safety.pptx')

    def test_focus_slide_uses_theme_and_other_hazards_keep_two_distinct_photo_pages(self):
        import re
        with tempfile.TemporaryDirectory() as root:
            Image.new('RGB', (600, 300), 'steelblue').save(Path(root)/'photo.jpg')
            rows = focus_fixture_rows() + [raw_issue(i, standard_detail_text='序号：50\n检查主题：其他对生产经营有影响的问题隐患\n评分说明：其他隐患规范') for i in range(20, 26)]
            report = fixture_report('photo.jpg', rows=rows, settings={'focus_category': '设备设施'})
            output = Path(root)/'safety.pptx'
            build_safety_presentation(report, output, root)
            prs = Presentation(output)
            text = '\n'.join(shape.text for shape in prs.slides[SLIDE_ORDER.index(12)].shapes if shape.has_text_frame)
            self.assertIn('本月重点工作类为设备设施，共4项', text)
            self.assertEqual(len(prs.slides), 31)
            photo_ids = []
            for index in (18, 19):
                page = prs.slides[SLIDE_ORDER.index(index)]
                self.assertEqual(sum(shape.shape_type == 13 and shape.top > 914400 for shape in page.shapes), 3)
                ids = list(map(int, re.findall(r'问题ID (\d+)', page.notes_slide.notes_text_frame.text)))
                self.assertEqual(len(ids), 3)
                photo_ids.extend(ids)
            self.assertEqual(sorted(photo_ids), list(range(20, 26)))
            focus_ids = set(map(int, re.findall(r'问题ID (\d+)', prs.slides[SLIDE_ORDER.index(12)].notes_slide.notes_text_frame.text)))
            self.assertTrue(focus_ids.issubset({1, 2, 3, 4}))


@unittest.skipUnless(os.environ.get('REPORT_AI_MEMORY_TEST_DSN'), 'requires isolated PostgreSQL test DSN')
class SafetySettingsPostgresTest(unittest.TestCase):
    def setUp(self):
        self.schema = 'safety_test_' + uuid.uuid4().hex
        self.dsn = os.environ['REPORT_AI_MEMORY_TEST_DSN']
        conn = psycopg2.connect(self.dsn); conn.autocommit = True
        with conn.cursor() as cur:
            cur.execute(f'CREATE SCHEMA "{self.schema}"')
            cur.execute(f'SET search_path TO "{self.schema}"')
            cur.execute('CREATE TABLE users(id INTEGER PRIMARY KEY); INSERT INTO users VALUES(7)')
            cur.execute('CREATE TABLE inspection_report_ai_memory(issue_ids BIGINT[], payload JSONB, created_at TIMESTAMPTZ, report_type TEXT, operation TEXT)')
        conn.close()
        self.engine = sa.create_engine('postgresql+psycopg2://', creator=lambda: psycopg2.connect(
            self.dsn, options=f'-csearch_path={self.schema}'))
        spec = importlib.util.spec_from_file_location('safety_migration', Path(__file__).parents[1]/'migrations/versions/20261029_001_safety_report.py')
        migration = importlib.util.module_from_spec(spec); spec.loader.exec_module(migration)
        with self.engine.begin() as connection:
            with Operations.context(MigrationContext.configure(connection)):
                migration.upgrade(); migration.upgrade()

    def connect(self):
        conn = psycopg2.connect(self.dsn, cursor_factory=RealDictCursor)
        with conn.cursor() as cur: cur.execute(f'SET search_path TO "{self.schema}"')
        conn.commit(); return conn

    def tearDown(self):
        self.engine.dispose()
        conn = psycopg2.connect(self.dsn); conn.autocommit = True
        with conn.cursor() as cur: cur.execute(f'DROP SCHEMA "{self.schema}" CASCADE')
        conn.close()

    def test_optimistic_revision_conflict_and_rollback_preserve_saved_settings(self):
        conn = self.connect()
        with conn.cursor() as cur: save_settings(cur, {'station_ids': [1]}, 0, {'id': 7, 'real_name': '甲'})
        conn.commit()
        with conn.cursor() as cur:
            with self.assertRaises(ValueError): save_settings(cur, {'station_ids': []}, 0, {'id': 7})
        conn.rollback()
        with conn.cursor() as cur:
            self.assertEqual(load_settings(cur)['settings'], {'station_ids': [1]})
            save_settings(cur, {'station_ids': []}, 1, {'id': 7})
        conn.rollback()
        with conn.cursor() as cur: self.assertEqual(load_settings(cur)['revision'], 1)
        conn.close()

    def test_two_concurrent_saves_cannot_overwrite_each_other(self):
        from concurrent.futures import ThreadPoolExecutor
        def save(value):
            conn = self.connect()
            try:
                with conn.cursor() as cur: save_settings(cur, {'station_ids': [value]}, 0, {'id': 7})
                conn.commit(); return True
            except ValueError:
                conn.rollback(); return False
            finally: conn.close()
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(save, [1, 2]))
        self.assertEqual(sorted(results), [False, True])
