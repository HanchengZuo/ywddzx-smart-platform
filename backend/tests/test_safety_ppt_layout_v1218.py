"""Safety PPT layout regressions; temporary files and mocked AI only."""
from copy import deepcopy
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from pptx import Presentation
from pptx.util import Inches

from report_ppt_artifacts import artifact_key
from safety_report_presentation import (TEMPLATE_FILE, build_safety_presentation,
    distribution_page, evidence_page, requirement_paragraph, suggestions_page)
from safety_template_styles import evidence_text_height
from tests.test_safety_report import fixture_report


class SafetyPptLayoutTest(unittest.TestCase):
    def test_requirements_deduplicate_in_photo_order_without_forced_line_breaks(self):
        def issue(text):return {'requirement':text,'requirement_status':'ready'}
        self.assertEqual(requirement_paragraph([issue('接地连接有效。')]*3),'规范要求：接地连接有效。')
        self.assertEqual(requirement_paragraph([issue('接地连接有效。'),issue(' 接地连接有效； '),issue('消防器材齐全。')]),
            '规范要求：1. 接地连接有效；2. 消防器材齐全。')
        self.assertEqual(requirement_paragraph([issue('接地连接有效。')]),'规范要求：接地连接有效。')

    def test_identical_requirements_leave_large_complete_photos_and_single_native_paragraph(self):
        from PIL import Image
        with tempfile.TemporaryDirectory() as root:
            root=Path(root);Image.new('RGB',(640,480),'steelblue').save(root/'photo.jpg')
            prs=Presentation(TEMPLATE_FILE)
            rows=[{'issue_id':i,'station_name':'示例站','description':'设施存在维护问题。','issue_photo':'photo.jpg',
                'requirement':'相关设施应保持完好、有效。','requirement_status':'manual'} for i in range(3)]
            for index in range(12,24):
                slide=prs.slides[index]
                evidence_page(slide,'分类问题',rows,{},rows,'实际巡查问题概述。',root)
                photos=[s for s in slide.shapes if s.name.startswith('safety-evidence-photo-')]
                self.assertEqual(len(photos),3)
                self.assertTrue(all(s.width>Inches(3.8) and s.height>Inches(2.8) for s in photos))
                self.assertTrue(all(s.crop_top==s.crop_bottom==s.crop_left==s.crop_right==0 for s in photos))
                norm=next(s for s in slide.shapes if s.name=='safety-evidence-requirements')
                self.assertEqual(len(norm.text_frame.paragraphs),1)
                self.assertEqual(norm.text,'规范要求：相关设施应保持完好、有效。')
                self.assertLessEqual(norm.top+norm.height,prs.slide_height)

    def test_pie_legends_are_centered_and_all_dense_percentages_remain_readable(self):
        prs=Presentation(TEMPLATE_FILE)
        report=fixture_report()
        section=deepcopy(report['sections'][0])
        for count in (1,5,14,19):
            with self.subTest(categories=count):
                page=Presentation(TEMPLATE_FILE).slides[10]
                section['detail_distribution']=[{'name':f'现场检查类别{i+1}','count':i+1,
                    'percentage':(i+1)*100/sum(range(1,count+1))} for i in range(count)]
                graph=next(s for s in prs.slides[7].shapes if s.has_chart)
                distribution_page(page,section,['第一点未加标点','第二点已加句号。','第三点以分号结束；'],graph)
                chart=next(s for s in page.shapes if s.has_chart)
                self.assertTrue(chart.chart._chartSpace.xpath('.//c:pie3DChart'))
                text=next(s for s in page.shapes if s.has_text_frame and s.text.startswith('问题分布'))
                self.assertTrue(all(p.text.endswith('。') for p in text.text_frame.paragraphs if p.text.startswith('（')))
                keys=[s for s in page.shapes if s.name=='safety-pie-key']
                labels=[s for s in page.shapes if s.name=='safety-pie-label']
                heading=next(s for s in page.shapes if s.name=='safety-pie-title')
                for top in {s.top for s in labels}:
                    row=[s for s in labels if s.top==top]
                    start=min(s.left for s in keys if abs(s.top-top-Inches(.07))<2)
                    end=max(s.left+s.width for s in row)
                    self.assertAlmostEqual((start+end)/2,heading.left+heading.width/2,delta=3)
                if count>10:
                    callouts=[s for s in page.shapes if s.name.startswith('safety-pie-percentage-')]
                    self.assertEqual(len(callouts),count)
                    self.assertEqual({s.text for s in callouts},{f'{(i+1)*100/sum(range(1,count+1)):.2f}%' for i in range(count)})
                    for left in {s.left for s in callouts}:
                        column=sorted((s for s in callouts if s.left==left),key=lambda s:s.top)
                        self.assertTrue(all(a.top+a.height<=b.top for a,b in zip(column,column[1:])))
                else:
                    root=chart.chart._chartSpace
                    self.assertEqual(root.xpath('.//c:pie3DChart/c:dLbls/c:showPercent')[0].get('val'),'1')
                    self.assertEqual(root.xpath('.//c:dLbls/c:spPr/a:solidFill/a:srgbClr')[0].get('val'),'FFFFFF')
                    self.assertFalse(root.xpath('.//c:dLbl/c:delete[@val="1"]'))

    def test_suggestions_keep_template_frames_red_squares_fixed_heading_and_complete_ai_text(self):
        for rectification,index in ((False,25),(True,29)):
            for length in (90,300):
                with self.subTest(rectification=rectification,length=length):
                    prs=Presentation(TEMPLATE_FILE);slide=prs.slides[index]
                    original=max((s for s in slide.shapes if s.has_text_frame and s.top>Inches(1)),key=lambda s:s.width*s.height)
                    box=(original.left,original.top,original.width,original.height)
                    items=[{'title':f'第{i+1}项管理建议','content':('应落实责任并及时整改。'*50)[:length]} for i in range(4)]
                    suggestions_page(slide,items,rectification=rectification)
                    body=next(s for s in slide.shapes if s.name.startswith('safety-') and s.has_text_frame)
                    self.assertEqual((body.left,body.top,body.width,body.height),box)
                    self.assertLessEqual(body.top+body.height,prs.slide_height)
                    self.assertEqual(len(body._element.xpath('.//a:buChar[@char="■"]')),5 if rectification else 4)
                    self.assertTrue(all(n.get('val')=='FF0000' for n in body._element.xpath('.//a:buClr/a:srgbClr')))
                    if rectification:self.assertEqual(body.text_frame.paragraphs[0].text,'针对问题整改的建议')
                    for item in items:self.assertIn(item['content'],body.text)
                    self.assertNotIn('…',body.text)
                    estimated=sum(evidence_text_height([p.text],body.width/914400-.55,p.runs[0].font.size.pt,
                        line_height=p.line_spacing.pt)-.1 for p in body.text_frame.paragraphs)
                    self.assertLessEqual(estimated+.12,body.height/914400)
        with self.assertRaisesRegex(ValueError,'工作建议文字过长'):
            suggestions_page(Presentation(TEMPLATE_FILE).slides[25],[{'title':'建议','content':'文'*12000}])

    def test_chapter_three_moves_as_one_block_after_all_video_evidence_and_cache_changes(self):
        with tempfile.TemporaryDirectory() as root:
            report=fixture_report();output=Path(root)/'report.pptx'
            with patch('safety_report_analysis.generate_analysis',side_effect=AssertionError('paid AI')):
                build_safety_presentation(report,output,root)
            prs=Presentation(output);self.assertEqual(len(prs.slides),31)
            text=lambda index:'\n'.join(s.text for s in prs.slides[index].shapes if s.has_text_frame)
            self.assertIn('1.月度重点工作类问题',text(9))
            self.assertIn('10. 油品质量控制类问题',text(20))
            self.assertEqual([s.text for s in prs.slides[21].shapes if s.has_text_frame],
                [s.text for s in Presentation(TEMPLATE_FILE).slides[9].shapes if s.has_text_frame])
            self.assertIn('三、问题分析-现场重点问题',text(22))
            self.assertIn('三、问题分析-视频重点问题',text(23))
            key=artifact_key('safety_quality',report,root)
            with patch('safety_report_presentation.RENDERER_VERSION','safety-native-1214-1'):
                self.assertNotEqual(artifact_key('safety_quality',report,root),key)
