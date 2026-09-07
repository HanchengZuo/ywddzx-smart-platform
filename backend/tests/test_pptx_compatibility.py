import hashlib
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch, MagicMock
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from pptx_compatibility import normalize_presentation, COMPATIBILITY_VERSION, DATA_TABLE_PREFIX
from report_ppt_artifacts import build_artifact, attach_export, export_preview_manifest, preview_slide_path


class PptCompatibilityTests(unittest.TestCase):
    def deck(self, values=(0,3,0), second=(0,0,0)):
        prs = Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        data = ChartData(); data.categories=['崇明第二','崇明第一','崇明第三']
        data.add_series('便利店卫生',values); data.add_series('商品摆放',second)
        shape = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(.5),Inches(1),Inches(12),Inches(5),data)
        return prs,slide,shape

    def test_zeroes_are_real_editable_cells_and_workbook_values(self):
        prs,slide,shape = self.deck()
        normalize_presentation(prs)
        table = next(s for s in slide.shapes if s.has_table)
        self.assertEqual(table.table.cell(0,1).text,'崇明第二')
        self.assertEqual(table.table.cell(1,1).text,'0')
        self.assertEqual(table.table.cell(1,2).text,'3')
        self.assertEqual(table.table.cell(2,3).text,'0')
        self.assertEqual(list(shape.chart.series[0].values),[0,3,0])
        self.assertTrue(shape.chart.part.chart_workbook.xlsx_part.blob)
        self.assertFalse(shape.chart._chartSpace.xpath('.//c:dTable'))
        self.assertLess(shape.top+shape.height,table.top)
        self.assertLessEqual(table.top+table.height,Inches(6))
        with tempfile.TemporaryDirectory() as root:
            path=Path(root)/'zero.pptx'; prs.save(path)
            restored=Presentation(path)
            self.assertEqual(restored.core_properties.version,COMPATIBILITY_VERSION)
            for node in restored.slides[0].shapes[0].chart._chartSpace.xpath('.//c:axId | .//c:crossAx'):
                self.assertGreaterEqual(int(node.get('val')),0)
            with ZipFile(path) as package:
                self.assertTrue(any(name.startswith('ppt/embeddings/') for name in package.namelist()))
                chart=package.read('ppt/charts/chart1.xml')
                self.assertIn(b'<c:v>0',chart)

    def test_all_zero_axis_and_idempotence(self):
        prs,slide,shape=self.deck((0,0,0))
        normalize_presentation(prs)
        geometry=[(s.left,s.top,s.width,s.height) for s in slide.shapes]
        normalize_presentation(prs)
        self.assertEqual(geometry,[(s.left,s.top,s.width,s.height) for s in slide.shapes])
        self.assertEqual(shape.chart.value_axis.maximum_scale,1)
        self.assertEqual(len([s for s in slide.shapes if s.name.startswith(DATA_TABLE_PREFIX)]),1)

    def test_decimal_values_are_not_rounded_to_zero(self):
        prs,slide,_=self.deck((0,.25,3.5))
        normalize_presentation(prs)
        table=next(s.table for s in slide.shapes if s.has_table)
        self.assertEqual([table.cell(1,n).text for n in (1,2,3)],['0','0.25','3.5'])

    def test_dense_data_uses_readable_continuation_without_clipping(self):
        prs, slide, shape = self.deck()
        shape.height = Inches(.7)
        normalize_presentation(prs)
        self.assertEqual(len(prs.slides), 2)
        table = next(s for s in prs.slides[1].shapes if s.has_table)
        self.assertEqual(table.table.cell(0, 1).text, '崇明第二')
        self.assertEqual(table.table.cell(1, 1).text, '0')
        self.assertLessEqual(table.top + table.height, prs.slide_height)
        normalize_presentation(prs)
        self.assertEqual(len(prs.slides), 2)

    def test_six_report_types_share_cached_ppt_and_previews(self):
        from tests.test_non_oil_report_presentation import NonOilReportPresentationTest
        report=NonOilReportPresentationTest().make_report()
        def render(path,directory):
            paths=[]
            for index in range(1,len(Presentation(path).slides)+1):
                target=Path(directory)/f'slide-{index:02d}.jpg'
                Image.new('RGB',(64,36),'white').save(target); paths.append(str(target))
            return paths
        with tempfile.TemporaryDirectory() as root, patch('non_oil_report_presentation._render_presentation_preview',side_effect=render) as renderer:
            for kind in ('quality_measurement','safety_quality','finance','on_site_service','equipment_facilities','non_oil'):
                with self.subTest(kind=kind):
                    directory,manifest=build_artifact(kind,report,root)
                    calls=renderer.call_count
                    reused,again=build_artifact(kind,report,root)
                    self.assertEqual(renderer.call_count,calls)
                    self.assertEqual((reused,again),(directory,manifest))
                    exported=Path(root)/(kind+'.pptx'); attach_export(directory,manifest,exported)
                    self.assertEqual(hashlib.sha256(exported.read_bytes()).hexdigest(),manifest['sha256'])
                    self.assertEqual(exported.read_bytes(),(directory/'report.pptx').read_bytes())
                    self.assertEqual(export_preview_manifest(exported,root),manifest)
                    self.assertTrue(preview_slide_path(exported,root,1).exists())
                    self.assertIsNone(preview_slide_path(exported,root,manifest['slide_count']+1))
                    self.assertIsNone(preview_slide_path(exported,root,0))
                    self.assertTrue(any(s.has_text_frame for p in Presentation(exported).slides for s in p.shapes))

    def test_cache_changes_with_payload_and_rejects_unsafe_sidecar(self):
        from report_ppt_artifacts import artifact_key
        import json
        with tempfile.TemporaryDirectory() as root:
            self.assertNotEqual(artifact_key('finance',{'n':1},root),artifact_key('finance',{'n':2},root))
            path=Path(root)/'export.pptx'
            path.with_suffix('.preview.json').write_text(json.dumps({'version':COMPATIBILITY_VERSION,'artifact_key':'../../etc'}))
            self.assertIsNone(export_preview_manifest(path,root))

    def test_expired_derived_cache_cleanup_preserves_saved_reports(self):
        import os
        from report_ppt_artifacts import cleanup_artifacts
        with tempfile.TemporaryDirectory() as root:
            base = Path(root) / 'report_presentations'; base.mkdir()
            stale = base / ('compatible-' + 'a' * 64); stale.mkdir()
            saved = base / 'saved-report'; saved.mkdir()
            recent = base / ('compatible-' + 'b' * 64); recent.mkdir()
            os.utime(stale, (0, 0)); os.utime(saved, (0, 0))
            cleanup_artifacts(root)
            self.assertFalse(stale.exists())
            self.assertTrue(saved.exists())
            self.assertTrue(recent.exists())

    def test_preview_requires_login(self):
        import app as reports
        with reports.app.test_request_context('/api/inspection-reports/exports/example/slides/1'):
            response, status = reports.require_signed_api_token()
        self.assertEqual(status, 401)

    def test_preview_permission_owner_and_missing_page(self):
        import app as reports
        for allowed in (False, True):
            with self.subTest(allowed=allowed), reports.app.test_request_context(
                '/api/inspection-reports/exports/example/slides/999?requested_by=999'
            ), patch('app.get_db_connection', return_value=MagicMock()), patch(
                'app.get_authorized_inspection_report_user',
                side_effect=None if allowed else PermissionError(), return_value={'id': 3}
            ), patch('app.get_inspection_report_export', return_value={
                'status': 'completed', 'file_path': 'report_exports/test.pptx'
            }) as lookup, patch('app.preview_slide_path', return_value=None):
                response, status = reports.get_report_ppt_preview_slide('example', 999)
            self.assertEqual(status, 404 if allowed else 403)
            if allowed:
                self.assertEqual(lookup.call_args.args[1:], ('example', 3))
            else:
                lookup.assert_not_called()

    def test_preview_returns_private_image_without_internal_errors(self):
        import app as reports
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / 'slide.jpg'; Image.new('RGB', (64, 36), 'white').save(path)
            with reports.app.test_request_context(), patch('app.get_db_connection', return_value=MagicMock()), patch(
                'app.get_authorized_inspection_report_user', return_value={'id': 3}
            ), patch('app.get_inspection_report_export', return_value={'status': 'completed'}), patch(
                'app.preview_slide_path', return_value=path
            ):
                response = reports.get_report_ppt_preview_slide('example', 1)
                self.assertEqual(response.mimetype, 'image/jpeg')
                self.assertEqual(response.headers['Cache-Control'], 'private, no-store')
                response.close()
            with reports.app.test_request_context(), patch('app.get_db_connection', side_effect=RuntimeError('internal secret')):
                response, status = reports.get_report_ppt_preview_slide('example', 1)
                self.assertEqual(status, 500)
                self.assertNotIn('internal secret', str(response.get_json()))
