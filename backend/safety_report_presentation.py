"""Editable safety template. Statistics and photos come only from the saved snapshot."""
from collections import Counter
from io import BytesIO
import math
from pathlib import Path
import re
import unicodedata
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx_compatibility import _normalize_package_fonts
from safety_report_data import ONSITE_PAGES, VIDEO_PAGES, unit_key, unit_label, category_rows
from safety_report_analysis import photo_examples
from safety_template_styles import (paragraph, fit_size, update_chart, frame_props, split_arrow,
    clone_pie_chart, compact_pie_layout, evidence_textbox, evidence_text_height, fit_evidence_layout,
    EVIDENCE_INTRO_BOX, EVIDENCE_CAPTION_BOXES, EVIDENCE_REQUIREMENTS_BOX)
from safety_video_themes import selected_video_group

TEMPLATE_FILE = Path(__file__).parent / 'assets/safety_report_template/template.pptx'
RENDERER_VERSION = 'safety-native-1218-1'
# Keep template indexes while composing, then move chapter three behind all
# chapter-two evidence. Only the slide ID order changes, not relationships.
SLIDE_ORDER = (*range(9), *range(12,24), 9, 10, 11, *range(24,31))
FONT = 'Microsoft YaHei'


def remove(shape):
    shape._element.getparent().remove(shape._element)


def text_units(text):
    return sum(1 if unicodedata.east_asian_width(c) in 'WF' else .55 for c in str(text))


def textbox(slide, text, x, y, w, h, size=18, color='111111', bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(.025)
    frame.margin_top = frame.margin_bottom = Inches(.02)
    # Bound text before shrinking. Keep the complete source in notes if a legacy
    # scoring clause cannot fit its fixed template box.
    text = str(text or '暂无数据。')
    selected = size
    while selected > 10:
        lines = sum(max(1, math.ceil(text_units(line) / max(1, w * 72 / selected))) for line in text.splitlines())
        if lines * selected * 1.18 <= h * 72 - 4: break
        selected -= 1
    max_chars = max(16, int(w * 72 / selected) * max(1, int((h * 72 - 4) / (selected * 1.18))))
    if text_units(text) > max_chars:
        slide.notes_slide.notes_text_frame.text += '\n完整文字：\n' + text
        cut = max(8, int(max_chars))
        text = text[:cut-1] + '…'
    for index, line in enumerate(text.splitlines()):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = line; p.alignment = align; p.space_after = Pt(2); p.line_spacing = 1.12
        for run in p.runs:
            run.font.name = FONT; run.font.size = Pt(selected); run.font.bold = bold
            run.font.color.rgb = RGBColor.from_string(color)
    return shape


def clear_dynamic(slide):
    for shape in list(slide.shapes):
        is_bar = shape.has_text_frame and bool(re.match(r'^\s*\d+[.、]', shape.text)) and shape.top < Inches(1.6)
        if shape.top >= Inches(.99) and not is_bar: remove(shape)
    used = {v for node in slide._element.iter() for k, v in node.attrib.items()
            if k.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}')}
    for rel in list(slide.part.rels.values()):
        if rel.reltype.endswith(('/image', '/chart', '/comments')) and rel.rId not in used:
            slide.part.drop_rel(rel.rId)


def set_bar(slide, title):
    for shape in slide.shapes:
        if shape.has_text_frame and shape.top < Inches(1.6) and re.match(r'^\s*\d+[.、]', shape.text):
            shape.text_frame.text = title
            for p in shape.text_frame.paragraphs:
                for run in p.runs:
                    run.font.name = FONT; run.font.size = Pt(22); run.font.bold = True
                    run.font.color.rgb = RGBColor(255,255,255)


def photo(slide, issue, root, x, y, w, h, *, bottom=False):
    raw = str(issue.get('issue_photo') or '').lstrip('/')
    path = (root / raw).resolve()
    if raw and path.is_relative_to(root.resolve()) and path.is_file():
        try:
            with Image.open(path) as original:
                image = ImageOps.exif_transpose(original).convert('RGB')
                width,height=image.size
                output=BytesIO();image.save(output,format='JPEG',quality=92);output.seek(0)
            scale=min(w/width,h/height)
            return slide.shapes.add_picture(output, Inches(x+(w-width*scale)/2), Inches(y+(h-height*scale)*(1 if bottom else .5)),
                width=Inches(width*scale),height=Inches(height*scale))
        except (OSError,ValueError): pass
    return textbox(slide,'暂无可用问题照片',x,y+h/2-.2,w,.4,15,'64748B',align=PP_ALIGN.CENTER)


def overview(slide, section):
    label='（一）四不两直现场检查方面' if section['mode']=='onsite' else '（二）视频扫站方面'
    body=next(s for s in slide.shapes if s.has_text_frame and s.text.startswith(('（一）','（二）')))
    graph=next(s for s in slide.shapes if s.has_chart)
    heading_props, body_props=frame_props(body), frame_props(body,1)
    body.height=graph.top-Inches(.12)-body.top
    units=sorted(section['units'],key=lambda r:unit_key(r['unit_name']))
    details='；'.join(f"{unit_label(r['unit_name'])}{r['station_count']}个站点，共{r['issue_count']}项，占比{r['percentage']:.1f}%" for r in units)
    narrative=section['overview']+'各管理片区及合资公司问题数量如下：'+(details+'。' if details else '当前范围暂无问题记录。')
    size=fit_size([narrative],body.width/914400,body.height/914400,18,reserved=32)
    body.text_frame.clear();body.text_frame.word_wrap=True
    paragraph(body,[(label,'111111')],20,first=True,bold=True,props=heading_props)
    spans=[];offset=0
    for match in re.finditer(r'(?<=加油站)\d+|(?<=发现问题)\d+|，站均问题\d+(?:\.\d+)?个',narrative):
        spans.extend([(narrative[offset:match.start()],'111111'),(match[0],'FF0000')]);offset=match.end()
    spans.append((narrative[offset:],'111111'))
    p=paragraph(body,spans,size,bold=True,props=body_props)
    p.line_spacing=Pt(size*1.45)
    update_chart(graph,units,(('站点数量','station_count'),('问题数量','issue_count')))


def distribution_page(slide, section, causes=None, pie_source=None):
    if causes is None:
        return template_distribution_page(slide,section)
    body=next(s for s in slide.shapes if s.has_text_frame and '问题分布' in s.text)
    distribution_props, cause_props, item_props=frame_props(body,1),frame_props(body,2),frame_props(body,3)
    for shape in list(slide.shapes):
        if shape.has_chart:
            refs=shape._element.xpath('.//c:chart/@r:id')
            remove(shape)
            for rid in refs:slide.part.drop_rel(rid)
    body.left=Inches(.55);body.top=Inches(1.35);body.width=Inches(6.2);body.height=Inches(5.76)
    rows=section['detail_distribution'];mode='现场' if section['mode']=='onsite' else '视频'
    detail='、'.join(f"{r['name']}（{r['count']}项）" for r in rows)
    text=f'问题分布：如右图所示，{mode}检查问题主要集中在{detail}，共{len(rows)}个方面。' if rows else f'问题分布：当前范围暂无{mode}问题。'
    reasons=[f'（{i+1}）{sentence(value)}' for i,value in enumerate(causes)] or ['当前范围暂无可分析的原因。']
    size=10
    for candidate in range(16,9,-1):
        height=sum(evidence_text_height([value],5.65,candidate,line_height=candidate*1.5)-.1
            for value in [text,'产生原因：',*reasons])
        if height+30/72<=5.76:
            size=candidate;break
    line_height=Pt(size*1.5)
    body.text_frame.clear();body.text_frame.word_wrap=True
    p=paragraph(body,[(text,'000000')],size,first=True,bold=True,props=distribution_props,bullet='square')
    p.line_spacing=line_height;p.space_after=Pt(10)
    p=paragraph(body,[('产生原因：','000000')],size,bold=True,props=cause_props,bullet='square')
    p.line_spacing=line_height;p.space_after=Pt(4)
    for reason in reasons:
        p=paragraph(body,[(reason,'000000')],size,bold=True,props=item_props)
        p.line_spacing=line_height;p.space_after=Pt(5)
    graph=clone_pie_chart(slide,pie_source,6.85,1.4,6.25,5.76)
    update_chart(graph,rows,(('问题数量','count'),),pie=True,title=mode+'检查问题分布')
    compact_pie_layout(slide,graph,rows,mode+'检查问题分布')


def template_distribution_page(slide, section):
    body=next(s for s in slide.shapes if s.has_text_frame and '本月检查共发现' in s.text)
    graph=next(s for s in slide.shapes if s.has_chart)
    intro_props=frame_props(body,1);item_props=frame_props(body,2)
    body.left=Inches(.68);body.top=Inches(1.85);body.width=Inches(5.05);body.height=Inches(4.85)
    for shape in slide.shapes:
        if shape.name.startswith('矩形') and shape.top>Inches(1):shape.height=Inches(5.244)
    rows=section['category_distribution'];mode='现场' if section['mode']=='onsite' else '视频'
    intro=f"本月检查共发现{mode}问题{section['total_issue_count']}项："
    def category_name(name): return name.removesuffix('类').removesuffix('类问题')+'类问题'
    texts=[f"{category_name(r['name'])}：{r['count']}项，占比{r['percentage']:.2f}%；" for r in rows]
    size=fit_size([intro,*texts],4.7,4.85,17,leading=1.6,reserved=len(rows)*6+18)
    body.text_frame.clear();body.text_frame.word_wrap=True
    p=paragraph(body,[(intro,'111111')],size,first=True,props=intro_props);p.line_spacing=1.3
    for i,(row,text) in enumerate(zip(rows,texts)):
        name=category_name(row['name'])+'：'
        p=paragraph(body,[(name,'111111'),(text[len(name):],'C00000' if i<(2 if mode=='视频' else 1) else '111111')],size,props=item_props,bullet=True)
        p.line_spacing=1.3;p.space_after=Pt(8)
    graph.left=Inches(6.15);graph.top=Inches(1.28);graph.width=Inches(6.95);graph.height=Inches(5.95)
    update_chart(graph,rows,(('问题数量','count'),),pie=True,title=f'{mode}巡查问题汇总分析')
    compact_pie_layout(slide,graph,rows,f'{mode}巡查问题汇总分析')


def video_typical_narrative(report):
    chosen=selected_video_group(report['safety_issues'],report['safety_analysis'].get('video:__typical__',{}),
        report.get('safety_settings',{}).get('typical_video_category',''))
    label=report['month_label']
    match=re.fullmatch(r'\d{4}年0?(\d+)月',label)
    if match:label=str(int(match[1]))+'月'
    if not chosen:return f'{label}视频扫站巡查过程中，暂无可统计的典型问题。'
    description=chosen.get('summary') or chosen['name']
    return f"{label}视频扫站巡查过程中，发现{description.rstrip('。；')}，总计出现{len(chosen['issue_ids'])}次。"


def typical_page(slide,report,root):
    findings={item['mode']:item for item in report['typical_findings']}
    for shape in list(slide.shapes):
        if shape.shape_type==13:remove(shape)
    for mode in ('onsite','video'):
        finding=findings.get(mode,{});issue=(finding.get('issues') or [{}])[0]
        left=mode=='onsite';x=.95 if left else 7.0
        if not left:
            chosen=selected_video_group(report['safety_issues'],report['safety_analysis'].get('video:__typical__',{}),
                report.get('safety_settings',{}).get('typical_video_category',''))
            if not chosen:issue={}
            elif issue.get('issue_id') not in chosen['issue_ids']:
                examples=photo_examples(report['safety_issues'],{'groups':[chosen]},1)
                issue=(examples or [{}])[0]
        title=next(s for s in slide.shapes if s.has_text_frame and s.text.startswith(('现场' if left else '视频')+'巡查典型问题'))
        props=frame_props(title);title.height=Inches(.5)
        for node in list(props) if props is not None else []:
            if node.tag.rsplit('}',1)[-1] in ('buFont','buChar','buNone'):props.remove(node)
        if props is not None:props.set('marL','0');props.set('indent','0')
        if props is not None:props.append(OxmlElement('a:buNone'))
        title.left=Inches(1.02 if left else 7.18);title.width=Inches(5.05)
        title.text_frame.clear()
        paragraph(title,[(('现场' if left else '视频')+'巡查典型问题：','000000')],20,first=True,bold=True,props=props)
        split_arrow(slide,.80 if left else 6.96,1.45)
        caption=next(s for s in slide.shapes if s.has_text_frame and s.name==('文本框 1' if left else '文本框 4'))
        caption.text_frame.clear()
        text=(finding.get('summary') or '暂无可用问题描述。') if left else video_typical_narrative(report)
        height=1.03
        while height<2.1 and fit_size([text],5.12,height,17,reserved=24 if left else 0)<14:
            height+=.12
        size=fit_size([text],5.12,height,17,reserved=24 if left else 0)
        caption.top=Inches(6.81-height);caption.width=Inches(5.12);caption.height=Inches(height)
        photo(slide,issue,root,x,1.95,5.15,(caption.top/914400)-2.01)
        if left:
            caption.left=Inches(.98)
            name=issue.get('station_name','')
            if name and not name.endswith('加油站'):name+='加油站'
            paragraph(caption,[(name or '暂无典型站点','000000')],17,first=True,bold=True,align=PP_ALIGN.CENTER)
            p=paragraph(caption,[(text,'000000')],size,bold=True);p.line_spacing=1.12
        else:
            caption.left=Inches(7.0)
            p=paragraph(caption,[(text,'000000')],size,first=True,bold=True);p.line_spacing=1.15


def evidence_layout(slide):
    def box(shape): return (shape.left, shape.top, shape.width, shape.height)
    captions=sorted((s for s in slide.shapes if s.has_text_frame
        and any(p.text.strip().endswith('加油站') for p in s.text_frame.paragraphs)
        and s.top>Inches(2)),key=lambda s:s.left)
    pictures=sorted((s for s in slide.shapes if s.shape_type==13 and s.top>Inches(1.6)),key=lambda s:s.left)
    requirement=next((s for s in slide.shapes if s.has_text_frame and s.text.startswith('规范要求')),None)
    intro=next((s for s in slide.shapes if s.has_text_frame and Inches(1.6)<s.top<Inches(2)),None)
    if len(captions)==len(pictures)==3:
        caption_boxes=[box(s) for s in captions];photo_boxes=[box(s) for s in pictures]
        station_first=captions[0].text_frame.paragraphs[0].text.strip().endswith('加油站')
    else:
        caption_boxes=list(EVIDENCE_CAPTION_BOXES)
        photo_boxes=[(x,Inches(2.79),w,Inches(2.2)) for x,_,w,_ in caption_boxes]
        station_first=False
    return (box(intro) if intro is not None else EVIDENCE_INTRO_BOX,caption_boxes,photo_boxes,
        box(requirement) if requirement is not None else EVIDENCE_REQUIREMENTS_BOX,
        requirement.line.width if requirement is not None else Pt(1),station_first,
        frame_props(intro) if intro is not None else None)


def confirmed_requirement(issue):
    text=str(issue.get('requirement') or '').strip()
    return text if issue.get('requirement_status') in ('ready','manual') else ''


def sentence(text):
    return str(text or '').strip().rstrip('。.;；，,!?！？')+'。'


def requirement_paragraph(issues):
    unique={}
    for issue in issues:
        value=re.sub(r'\s+',' ',confirmed_requirement(issue)).strip() or '暂无已确认的规范要求。'
        key=re.sub(r'\s+','',value).rstrip('。.;；')
        unique.setdefault(key,value.rstrip('。.;；'))
    values=list(unique.values())
    if not values:return '规范要求：当前范围暂无可对应的已确认规范要求。'
    if len(values)==1:return '规范要求：'+values[0]+'。'
    return '规范要求：'+'；'.join(f'{index+1}. {text}' for index,text in enumerate(values))+'。'


def evidence_page(slide, title, issues, analysis, chosen, intro, root):
    intro_box,captions,pictures,norm_box,border,station_first,intro_props=evidence_layout(slide)
    lookup={row['issue_id']:row for row in issues}
    selected=[{**lookup[item['issue_id']],'caption':item.get('caption')}
        for item in chosen if item.get('issue_id') in lookup][:3]
    lines=[requirement_paragraph(selected)]
    descriptions=[]
    for issue in selected:
        station=str(issue.get('station_name') or '').strip()
        if station and not station.endswith('加油站'):station+='加油站'
        descriptions.append((issue.get('caption') or issue.get('description') or '暂无问题描述。',station or '暂无站点信息'))
    spacing=intro_props.find(f"{qn('a:lnSpc')}/{qn('a:spcPct')}") if intro_props is not None else None
    intro_leading=int(spacing.get('val'))/100000 if spacing is not None else 1.25
    layout=fit_evidence_layout(intro,descriptions,lines,intro_box,captions,pictures,norm_box,intro_leading)
    clear_dynamic(slide);set_bar(slide,title)
    heading=evidence_textbox(slide,layout['intro'],'safety-evidence-intro')
    p=paragraph(heading,[(intro,'262626')],layout['intro_size'],first=True,bold=True,props=intro_props)
    p.line_spacing=intro_leading
    if not selected:
        textbox(slide,'当前范围暂无可展示的问题证据。',.6,3.3,12.1,.6,18,'64748B',align=PP_ALIGN.CENTER)
    for i,issue in enumerate(selected):
        image=photo(slide,issue,root,*(v/914400 for v in layout['photos'][i]),bottom=True)
        image.name=f'safety-evidence-photo-{i+1}'
        caption=evidence_textbox(slide,layout['captions'][i],f'safety-evidence-caption-{i+1}')
        description,station=descriptions[i]
        spans=[(description,'000000',False),(station,'3F95D2',True)]
        if station_first:spans.reverse()
        for j,(text,color,bold) in enumerate(spans):
            p=paragraph(caption,[(text,color)],layout['caption_size'],first=j==0,bold=bold,align=PP_ALIGN.CENTER)
            p.line_spacing=1;p.space_after=Pt(6 if j==0 else 0)
        requirement=confirmed_requirement(issue) or '暂无已确认的规范要求。'
        slide.notes_slide.notes_text_frame.text += (f"\n照片{i+1}：问题ID {issue['issue_id']}，{station}"
            f"\n问题原文：{issue.get('description') or ''}\n规范要求：{requirement}\n")
    norm=evidence_textbox(slide,layout['norm'],'safety-evidence-requirements',border=border)
    for i,text in enumerate(lines):
        p=paragraph(norm,[(text,'C00000')],layout['norm_size'],first=i==0,bold=True,align=PP_ALIGN.LEFT)
        p.line_spacing=layout['norm_leading']


def suggestions_page(slide, items, *, rectification=False):
    """Retain native headings/indentation, fitting complete AI text before drawing."""
    body=max((s for s in slide.shapes if s.has_text_frame and s.top>Inches(1)),
        key=lambda s:s.width*s.height)
    heading_props=frame_props(body)
    item_props=frame_props(body,1 if rectification else 0)
    content_props=frame_props(body,2 if rectification else 1)
    entries=[]
    if rectification:entries.append(('针对问题整改的建议',True,heading_props))
    for index,item in enumerate(items):
        title=str(item.get('title') or '').strip()
        title=re.sub(r'^(?:[一二三四五六七八九十]+、|\d+[、.])\s*','',title)
        heading=title if rectification else f'{"一二三四"[index]}、{title}'
        entries.append((heading,True,item_props))
        entries.extend((line.strip(),False,content_props) for line in str(item.get('content') or '').splitlines() if line.strip())
    if not items:entries.append(('当前范围暂无对应问题，无需新增建议。',False,content_props))
    width=body.width/914400-.55
    available=body.height/914400
    selected=None
    for size in range(16,9,-1):
        for leading in (1.5,1.35,1.2,1.1):
            line_height=size*leading
            height=sum(evidence_text_height([text],width,size,line_height=line_height)-.1 for text,_,_ in entries)
            if height+.12<=available:
                selected=(size,line_height);break
        if selected:break
    if not selected:raise ValueError('工作建议文字过长，无法完整放入模板，请精简建议后重试。')
    size,line_height=selected
    body.name='safety-rectification-suggestions' if rectification else 'safety-work-suggestions'
    body.text_frame.clear();body.text_frame.word_wrap=True
    from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR
    body.text_frame.auto_size=MSO_AUTO_SIZE.NONE
    body.text_frame.vertical_anchor=MSO_ANCHOR.TOP
    for index,(text,heading,props) in enumerate(entries):
        p=paragraph(body,[(text,'262626')],size,first=index==0,bold=heading,props=props,bullet='square' if heading else False)
        p.line_spacing=Pt(line_height)


def backlog_page(slide, rows, mode, cutoff):
    body=next(s for s in slide.shapes if s.has_text_frame and s.text.startswith(('（一）','（二）')))
    graph=next(s for s in slide.shapes if s.has_chart)
    heading_props,body_props=frame_props(body),frame_props(body,1)
    body.height=graph.top-Inches(.12)-body.top
    onsite=mode=='onsite';name='现场' if onsite else '视频'
    heading=[(('（一）' if onsite else '（二）')+'各管理片区及合资公司','262626'),(name,'FF0000'),
        (('检查' if onsite else '')+f'问题整改情况（数据截至{cutoff}）','262626')]
    narrative='整改情况如下：'+('；'.join(f"{unit_label(r['unit_name'])}待签名{int(r.get('pending_sign') or 0)}项、"
        f"待整改{int(r.get('pending_rectification') or 0)}项" for r in rows)+'。' if rows else '当前范围暂无待签名或待整改问题。')
    # Preserve the source's first-line indent and fixed chart rectangle. Dense
    # unit lists compact only the narrative, without clipping or extra tables.
    width=body.width/914400-.59375
    heading_height=evidence_text_height([''.join(text for text,_ in heading)],width,16,line_height=28)
    size=line_height=None
    for candidate in range(16,9,-1):
        for factor in (1.75,1.5,1.35,1.25):
            height=evidence_text_height([narrative],width,candidate,line_height=candidate*factor)
            if heading_height+height-.1<=body.height/914400:
                size=candidate;line_height=candidate*factor;break
        if size is not None:break
    if size is None:raise ValueError('整改统计文字过长，无法完整放入模板的原生图表上方。')
    body.text_frame.clear();body.text_frame.word_wrap=True
    p=paragraph(body,heading,16,first=True,props=heading_props);p.line_spacing=Pt(28)
    spans=[];offset=0
    for match in re.finditer(r'(?<=待签名)\d+|(?<=待整改)\d+',narrative):
        spans.extend([(narrative[offset:match.start()],'262626'),(match[0],'FF0000')]);offset=match.end()
    spans.append((narrative[offset:],'262626'))
    p=paragraph(body,spans,size,props=body_props);p.line_spacing=Pt(line_height)
    update_chart(graph,rows,(('待签名','pending_sign'),('待整改','pending_rectification')),preserve_style=True)


def build_safety_presentation(report, output_path, storage_root=None):
    if report.get('safety_template_version') != 2:
        raise ValueError('安全质量报告已更换模板，请先重新生成报告。')
    prs=Presentation(TEMPLATE_FILE); source=list(prs.slides); root=Path(storage_root or '.')
    if len(source)!=31: raise ValueError('安全报告模板页数异常。')
    for slide in source:
        slide.notes_slide.notes_text_frame.text=''
        for rel in list(slide.part.rels.values()):
            if rel.reltype.endswith('/comments'): slide.part.drop_rel(rel.rId)
    label=report['month_label']
    for shape in source[0].shapes:
        if shape.has_text_frame and re.search(r'\d{4}年\s*\d+月',shape.text):
            for index, paragraph in enumerate(shape.text_frame.paragraphs):
                for j, run in enumerate(paragraph.runs):
                    run.text = label if index == 0 and j == 0 else ''
                    # Retain the template's explicit color, including its hidden date placeholder.
                    if run.text and len(label) > 12: run.font.size = Pt(20)
    sections={s['mode']:s for s in report['sections']}
    overview(source[2],sections['onsite']);overview(source[3],sections['video'])
    typical_page(source[5],report,root)
    distribution_page(source[7],sections['onsite']);distribution_page(source[8],sections['video'])
    distribution_page(source[10],sections['onsite'],report['analysis_causes']['onsite'],next(s for s in source[7].shapes if s.has_chart))
    distribution_page(source[11],sections['video'],report['analysis_causes']['video'],next(s for s in source[8].shapes if s.has_chart))
    issues=report['safety_issues'];analyses=report['safety_analysis'];focus=report['focus_category']
    key='onsite:topic:'+focus;value=analyses.get(key,{})
    selected=category_rows(issues,'onsite','topic:'+focus)
    names='、'.join(g['name'] for g in value.get('groups',[])[:3])
    evidence_page(source[12],'1.月度重点工作类问题',selected,value,photo_examples(selected,value,3,distinct_groups=True),
        f'本月重点工作类为{focus or "暂无"}，共{len(selected)}项。主要问题集中在{names or "暂无对应问题"}方面。',root)
    targets=[(13,'设备设施',0),(14,'设备设施',3),(15,'污染防治',0),(16,'能力培训和意识',0),(17,'风险合规',0),(18,'其他隐患类',0),(19,'其他隐患类',3)]
    titles={key:title for key,title,_ in ONSITE_PAGES}
    for page,category,offset in targets:
        selected=category_rows(issues,'onsite',category);value=analyses.get('onsite:'+category,{})
        names='、'.join(g['name'] for g in value.get('groups',[])[:4])
        evidence_page(source[page],titles[category],selected,value,photo_examples(selected,value,offset+3)[offset:],
            f'本月{category}问题共{len(selected)}项，其中问题主要集中在{names or "暂无对应问题"}。',root)
    for page,(category,title) in enumerate(VIDEO_PAGES,20):
        selected=category_rows(issues,'video',category);value=analyses.get('video:'+category,{})
        unit_counts=Counter(row['unit_name'] for row in selected)
        most=sorted(unit_counts.items(),key=lambda item:(-item[1],unit_key(item[0])))[:1]
        share=f"其中{most[0][0]}存在{most[0][1]}项，占比{most[0][1]/len(selected)*100:.1f}%。" if most else ''
        names='、'.join(g['name'] for g in value.get('groups',[])[:4])
        evidence_page(source[page],title,selected,value,photo_examples(selected,value,3),
            f'本月{category}问题共{len(selected)}项。{share}问题主要集中在{names or "暂无对应问题"}。',root)
    for page,key in [(25,'work_suggestions'),(29,'rectification_suggestions')]:
        suggestions_page(source[page],report.get(key) or [],rectification=page==29)
    cutoff=(report.get('summary') or {}).get('date_to') or label
    for page,mode in [(27,'onsite'),(28,'video')]:
        rows=sorted(report['pending_rectification'].get(mode) or [],key=lambda r:unit_key(r['unit_name']))
        backlog_page(source[page],rows,mode,cutoff)
    _normalize_package_fonts(prs)
    slide_ids=list(prs.slides._sldIdLst)
    for node in slide_ids:prs.slides._sldIdLst.remove(node)
    for index in SLIDE_ORDER:prs.slides._sldIdLst.append(slide_ids[index])
    output_path=Path(output_path);output_path.parent.mkdir(parents=True,exist_ok=True);prs.save(output_path)
    return {'ppt_path':str(output_path),'slide_count':len(prs.slides)}
